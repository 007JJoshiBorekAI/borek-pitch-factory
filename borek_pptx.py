"""
borek_pptx.py - Borek Solutions Group presentation engine (python-pptx)

Re-creates the layouts of "Borek Master Presentation.html" / "Ai Tech Borek Presentation EN.pptx"
1:1 on a 1920 x 1080 px grid (1 px = 9525 EMU = 0.75 pt), Inter font, official palette.

Two entry points use this module:
    make_ai_tech_deck.py   - text file  -> max 8 slides
    make_master_deck.py    - pitch + client + transcript + text -> deck built from the master layouts

A deck is a list of slide specs:  [{"layout": "contrast", ...fields...}, ...]
See LAYOUT_DOCS (bottom of this file) for the fields of every layout.
"""
from __future__ import annotations

import copy
import os
import re
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

HERE = Path(__file__).resolve().parent
ASSETS = HERE / "borek_assets"

# --------------------------------------------------------------------------------------
# Design tokens (Master guide, slide 03)
# --------------------------------------------------------------------------------------
PX = 9525
SLIDE_W, SLIDE_H = 1920, 1080
NAVY, SLATE, GREY, LGREY = "0D1240", "515C70", "8A90A5", "9AA0B3"
PANEL, BORDER, WHITE = "F3F4F8", "E4E7F0", "FFFFFF"
BLUE, TEAL, RED, ORANGE = "124F94", "02A69F", "DD3D00", "E07E00"
ACCENTS = [BLUE, TEAL, RED, ORANGE]
RULE = "D5DDE9"          # thin rule inside cards
BODY_ALT = "5C6178"      # body text of the four-card layout
FONT = "Inter"
NATURAL_LH = 1.2083      # Inter natural line height (ascent+descent)

# Inter weights: the reference PPTX uses plain "Inter" + bold flag. Set BOREK_FONT_VARIANTS=1 to use
# the static family names "Inter Light / Inter Medium / Inter SemiBold" instead (needs those installed).
FONT_VARIANTS = os.environ.get("BOREK_FONT_VARIANTS", "0") == "1"


def font_spec(weight: int):
    if FONT_VARIANTS:
        name = {300: "Inter Light", 500: "Inter Medium", 600: "Inter SemiBold"}.get(weight, FONT)
        return name, weight >= 700
    return FONT, weight >= 600


# --------------------------------------------------------------------------------------
# Brand data: official roster (Master guide, slide 04) - names / positions must not change
# --------------------------------------------------------------------------------------
ROSTER = {
    "Konstantin Borek": dict(role="Group CEO", photo="konstantin_borek", email="borek@boreksolutions.de",
                             phone="+49 151 539 652 10"),
    "Fiona Oldenburg": dict(role="Group COO", photo="fiona_oldenburg", email="fiona.oldenburg@boreksolutions.de",
                            phone="+49 151 729 76084"),
    "Kushal Rao": dict(role="Co-Founder & CEO India", photo="kushal_rao"),
    "Viktoria Schünemann": dict(role="Managing Partner & Co-CCO", role_long="Managing Partner & Co-Chief Commercial Officer",
                                photo="viktoria_schuenemann", email="viktoria.schuenemann@boreksolutions.de",
                                phone="+49 151 223 06 286"),
    "Katharina Bahlsen": dict(role="Managing Partner AI Tech", photo="katharina_bahlsen"),
    "Elena Manovska": dict(role="Managing Partner & Co-CCO", role_long="Managing Partner & Co-Chief Commercial Officer",
                           photo="elena_manovska", email="manovska@boreksolutions.de", phone="+389 71 212 702"),
    "Muhamet Abdullahu": dict(role="CEO Kosovo", photo="muhamet_abdullahu"),
    "Max Fichtner": dict(role="Senior Strategic Partnership Manager", photo="max_fichtner",
                         email="max.fichtner@boreksolutions.de", phone="+49 171 9994310"),
    "Meena Dholakia": dict(role="Vice President Customer Success", photo="meena_dholakia"),
}
F2_TEAM = ["Konstantin Borek", "Fiona Oldenburg", "Kushal Rao", "Viktoria Schünemann", "Katharina Bahlsen",
           "Elena Manovska", "Muhamet Abdullahu"]

# Fixed "Who we are" slide (F2) - reused unchanged in every deck
F2_TIMELINE_MASTER = [
    ("1781", "Company founded"), ("1790", "Bookbindery in Bodenwerder"), ("1920", "Forms printing begins"),
    ("1965", "Leading forms printer in northern Germany"), ("2014", "Borek IT Sourcing, India"),
    ("2016", "7th generation · ISO 27001"), ("2021", "Borek Solutions Group"), ("2022", "Hub Prishtina, Kosovo"),
    ("2023", "Focus on IT & AI services"), ("2026", "Borek AI Suite"),
]
F2_TIMELINE_AITECH = [
    ("1781", "Pigge Druck und Service founded by Wilhelm E. Pigge"),
    ("1790", "From newspaper printing to bookbinding in Bodenwerder"),
    ("1920", "Move into printing of commercial and insurance forms"),
    ("1965", "Leading forms printer in Northern Germany"),
    ("2014", "Borek IT Sourcing Private Limited founded in India as a subsidiary of the printing company"),
    ("2016", "Konstantin Borek (7th generation) joins · ISO 27001 for the Indian subsidiary"),
    ("2021", "Printing business closed — focus on expanding Borek Solutions Group"),
    ("2022", "Borek Solutions LLC founded in Prishtina, Kosovo"),
    ("2023", "Focus on IT & AI services"),
    ("2026", "Borek Solutions Bulgaria founded · launch of Borek AI Suite"),
]
LOCATIONS = [("Braunschweig", "Headquarter, Germany"), ("Prishtina", "AI-Hub"), ("Bulgaria", "AI-Hub"),
             ("Vadodara", "AI-Operations-Hub")]
ADDRESS = "Altewiekring 20A, 38102 Braunschweig, Germany"
HQ_PHONE = "+49 531 28354 162"


# --------------------------------------------------------------------------------------
# Text measurement (uses the bundled Inter variable font when Pillow can load it)
# --------------------------------------------------------------------------------------
_FONT_CACHE: dict = {}
_FONT_PATH = ASSETS / "fonts" / "Inter-Variable.ttf"
_MEASURE_OK = True


def _pil_font(size: float, weight: int):
    global _MEASURE_OK
    key = (round(size * 4), weight)
    if key in _FONT_CACHE:
        return _FONT_CACHE[key]
    f = None
    if _MEASURE_OK:
        try:
            from PIL import ImageFont
            f = ImageFont.truetype(str(_FONT_PATH), size)
            try:
                f.set_variation_by_axes([weight])
            except Exception:
                pass
        except Exception:
            _MEASURE_OK = False
            f = None
    _FONT_CACHE[key] = f
    return f


def eff_weight(weight: int) -> int:
    """Weight that PowerPoint will really render (plain 'Inter' only has regular + bold)."""
    if FONT_VARIANTS:
        return weight
    return 700 if weight >= 600 else 400


def text_w(text: str, size: float, weight: int = 400, ls: float = 0.0) -> float:
    weight = eff_weight(weight)
    f = _pil_font(size, weight)
    if f is not None:
        try:
            return f.getlength(text) + ls * len(text)
        except Exception:
            pass
    return len(text) * (size * (0.60 if weight >= 600 else 0.56) + ls)


def wrap_lines(text: str, size: float, weight: int, width: float, ls: float = 0.0) -> list[str]:
    out: list[str] = []
    for para in str(text).split("\n"):
        words = para.split(" ")
        cur = ""
        for w in words:
            trial = (cur + " " + w) if cur else w
            if cur and text_w(trial, size, weight, ls) > width:
                out.append(cur)
                cur = w
            else:
                cur = trial
        out.append(cur)
    return out


def th(text, size, weight=400, w=9999, lh=1.4, ls=0.0, caps=False) -> float:
    """Estimated height of a wrapped text block."""
    t = str(text).upper() if caps else str(text)
    return len(wrap_lines(t, size, weight, w, ls)) * size * lh


def nlines(text, size, weight, w, ls=0.0, caps=False) -> int:
    t = str(text).upper() if caps else str(text)
    return len(wrap_lines(t, size, weight, w, ls))


# --------------------------------------------------------------------------------------
# XML helpers
# --------------------------------------------------------------------------------------
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _alpha(clr_parent, alpha: float):
    """Append <a:alpha> to the srgbClr that lives inside clr_parent (a solidFill element)."""
    clr = clr_parent.find(qn("a:srgbClr"))
    if clr is not None:
        a = etree.SubElement(clr, qn("a:alpha"))
        a.set("val", str(int(round(alpha * 100000))))


def _strip_style(shape):
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def mix(fg: str, bg: str, a: float) -> str:
    f = [int(fg[i:i + 2], 16) for i in (0, 2, 4)]
    b = [int(bg[i:i + 2], 16) for i in (0, 2, 4)]
    return "".join(f"{round(f[i] * a + b[i] * (1 - a)):02X}" for i in range(3))


class Run:
    """Text run with optional style overrides: Run('text', weight=700, color=BLUE, size=30)."""

    def __init__(self, text, **kw):
        self.text, self.kw = text, kw


# --------------------------------------------------------------------------------------
# Slide wrapper
# --------------------------------------------------------------------------------------
class S:
    def __init__(self, deck: "Deck", slide, dark=False, footer="auto"):
        self.deck, self.slide, self.dark = deck, slide, dark
        self.footer = footer          # "auto" | "none" | "closing"
        self.max_bottom = 0.0
        self.warnings: list[str] = []
        self.label = ""

    # ---- shapes ---------------------------------------------------------------------
    def rect(self, x, y, w, h, fill=None, line=None, lw=1.0, r=0, fill_alpha=None, line_alpha=None,
             dash=False, oval=False):
        if line:   # CSS borders are inside the box, PowerPoint strokes straddle it
            x, y, w, h = x + lw / 2, y + lw / 2, w - lw, h - lw
        kind = MSO_SHAPE.OVAL if oval else (MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE)
        sp = self.slide.shapes.add_shape(kind, Emu(int(x * PX)), Emu(int(y * PX)), Emu(max(int(w * PX), 1)),
                                         Emu(max(int(h * PX), 1)))
        _strip_style(sp)
        if r and not oval:
            sp.adjustments[0] = min(0.5, r / max(min(w, h), 1))
        if fill:
            sp.fill.solid()
            sp.fill.fore_color.rgb = RGBColor.from_string(fill)
            if fill_alpha is not None:
                _alpha(sp._element.spPr.find(qn("a:solidFill")), fill_alpha)
        else:
            sp.fill.background()
        if line:
            sp.line.color.rgb = RGBColor.from_string(line)
            sp.line.width = Emu(int(lw * PX))
            if line_alpha is not None:
                _alpha(sp._element.spPr.find(qn("a:ln")).find(qn("a:solidFill")), line_alpha)
            if dash:
                sp.line.dash_style = MSO_LINE.DASH
        else:
            sp.line.fill.background()
        self._track(y + h)
        return sp

    def hline(self, x, y, w, color=BORDER, th_=1.0, alpha=None):
        return self.rect(x, y, w, th_, fill=color, fill_alpha=alpha)

    def gradient(self, x, y, w, h, stops, angle=0):
        """stops = [(pos 0..1, 'RRGGBB', alpha 0..1)]"""
        sp = self.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(int(x * PX)), Emu(int(y * PX)),
                                         Emu(int(w * PX)), Emu(int(h * PX)))
        _strip_style(sp)
        sp.line.fill.background()
        spPr = sp._element.spPr
        for tag in ("a:solidFill", "a:noFill", "a:gradFill"):
            e = spPr.find(qn(tag))
            if e is not None:
                spPr.remove(e)
        grad = etree.Element(qn("a:gradFill"), rotWithShape="1")
        gs_lst = etree.SubElement(grad, qn("a:gsLst"))
        for pos, col, al in stops:
            gs = etree.SubElement(gs_lst, qn("a:gs"), pos=str(int(pos * 100000)))
            c = etree.SubElement(gs, qn("a:srgbClr"), val=col)
            etree.SubElement(c, qn("a:alpha"), val=str(int(al * 100000)))
        etree.SubElement(grad, qn("a:lin"), ang=str(int(angle * 60000)), scaled="0")
        geom = spPr.find(qn("a:prstGeom"))
        geom.addnext(grad)
        return sp

    def pic(self, path, x, y, w, h=None, shape="rect", radius=0, alpha=None):
        from PIL import Image
        path = str(path)
        iw, ih = Image.open(path).size
        if h is None:
            h = w * ih / iw
        pic = self.slide.shapes.add_picture(path, Emu(int(x * PX)), Emu(int(y * PX)), Emu(int(w * PX)),
                                            Emu(int(h * PX)))
        ia, ba = iw / ih, w / h           # object-fit: cover
        if ia > ba + 1e-3:
            e = (1 - ba / ia) / 2
            pic.crop_left = pic.crop_right = e
        elif ia < ba - 1e-3:
            e = (1 - ia / ba) / 2
            pic.crop_top = pic.crop_bottom = e
        geom = pic._element.spPr.find(qn("a:prstGeom"))
        if shape == "ellipse":
            geom.set("prst", "ellipse")
        elif shape == "round":
            geom.set("prst", "roundRect")
            av = geom.find(qn("a:avLst"))
            if av is None:
                av = etree.SubElement(geom, qn("a:avLst"))
            gd = etree.SubElement(av, qn("a:gd"), name="adj", fmla=f"val {int(min(0.5, radius / min(w, h)) * 100000)}")
        if alpha is not None:
            blip = pic._element.blipFill.find(qn("a:blip"))
            etree.SubElement(blip, qn("a:alphaModFix"), amt=str(int(alpha * 100000)))
        self._track(y + h)
        return pic

    # ---- text -----------------------------------------------------------------------
    def text(self, x, y, w, content, size=21, weight=400, color=SLATE, lh=1.4, ls=0.0, align="l", caps=False,
             italic=False, h=None, wrap=True, anchor="t", name=None):
        """Add a text box. `content` is str ('\\n' = new paragraph) or list of str / Run.
        Returns the (estimated) bottom y of the text."""
        runs = content if isinstance(content, (list, tuple)) else [content]
        runs = [r if isinstance(r, Run) else Run(str(r)) for r in runs]
        if caps:
            for r in runs:
                r.text = r.text.upper()
        flat = "".join(r.text for r in runs)
        big = max([r.kw.get("size", size) for r in runs] + [size])
        if wrap:
            n = nlines(flat, size, weight, w, ls)
            if n * size * lh > (h or 1e9) + 1:
                self.warnings.append(f"text may overflow its box: '{flat[:40]}...'")
        else:
            n = flat.count("\n") + 1
        height = h if h is not None else n * big * lh + 2
        shift = -(lh - NATURAL_LH) * size / 2
        tb = self.slide.shapes.add_textbox(Emu(int(x * PX)), Emu(int((y + shift) * PX)), Emu(int(w * PX)),
                                           Emu(int(height * PX)))
        tf = tb.text_frame
        tf.word_wrap = wrap
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
        # split runs into paragraphs on "\n"
        paras: list[list[Run]] = [[]]
        for r in runs:
            parts = r.text.split("\n")
            for i, part in enumerate(parts):
                if i > 0:
                    paras.append([])
                if part != "" or len(parts) == 1:
                    paras[-1].append(Run(part, **r.kw))
        for i, pr in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
            p.line_spacing = lh / NATURAL_LH
            for r in pr:
                kw = dict(weight=weight, color=color, size=size, italic=italic, ls=ls)
                kw.update(r.kw)
                run = p.add_run()
                run.text = r.text
                f = run.font
                fname, bold = font_spec(kw["weight"])
                f.name = fname
                f.size = Pt(kw["size"] * 0.75)
                f.bold = bold
                f.italic = bool(kw["italic"])
                f.color.rgb = RGBColor.from_string(kw["color"])
                if kw["ls"]:
                    run._r.get_or_add_rPr().set("spc", str(int(round(kw["ls"] * 75))))
        bottom = y + (n * big * lh)
        self._track(bottom)
        if name:
            tb.name = name
        return bottom

    def _track(self, bottom):
        self.max_bottom = max(self.max_bottom, bottom)


# --------------------------------------------------------------------------------------
# Deck
# --------------------------------------------------------------------------------------
class Deck:
    def __init__(self, template: str | os.PathLike | None = None, footer_text="Borek Solutions Group · boreksolutions.de · Confidential"):
        if template and Path(template).exists():
            self.prs = Presentation(str(template))
            lst = self.prs.slides._sldIdLst
            for sld in list(lst):                     # keep theme / size / layout, drop the old slides
                self.prs.part.drop_rel(sld.rId)
                lst.remove(sld)
            self.layout = self.prs.slide_layouts[0]
        else:
            self.prs = Presentation()
            self.prs.slide_width, self.prs.slide_height = Emu(SLIDE_W * PX), Emu(SLIDE_H * PX)
            self.layout = self.prs.slide_layouts[6]
        self.slides: list[S] = []
        self.footer_text = footer_text
        self.total_override = None

    def new(self, dark=False, bg=None, footer="auto") -> S:
        sl = self.prs.slides.add_slide(self.layout)
        for ph in list(sl.placeholders):
            ph._element.getparent().remove(ph._element)
        sl.background.fill.solid()
        sl.background.fill.fore_color.rgb = RGBColor.from_string(bg or (NAVY if dark else WHITE))
        s = S(self, sl, dark, footer)
        self.slides.append(s)
        return s

    def _footers(self):
        total = self.total_override or len(self.slides)
        for i, s in enumerate(self.slides, 1):
            if s.footer == "none":
                continue
            num = f"{i:02d} / {total:02d}"
            if s.footer == "closing":
                s.text(90, 1036, 600, "boreksolutions.de", size=15, color=mix(WHITE, NAVY, 0.7), ls=2.6, caps=True,
                       lh=1.2, wrap=False)
                s.text(1830 - 300, 1036, 300, num, size=15, color=mix(WHITE, NAVY, 0.7), ls=2.6, lh=1.2, align="r",
                       wrap=False)
            elif s.dark:
                c = mix(WHITE, NAVY, 0.7)
                s.text(115, 1010, 1200, self.footer_text, size=15, color=c, ls=2.6, caps=True, lh=1.2, wrap=False)
                s.text(1805 - 300, 1010, 300, num, size=15, color=c, ls=2.6, lh=1.2, align="r", wrap=False)
            else:
                s.text(72, 1000, 1400, self.footer_text, size=15, color=LGREY, ls=2.6, caps=True, lh=1.2, wrap=False)
                s.text(1848 - 300, 1000, 300, num, size=15, color=LGREY, ls=2.6, lh=1.2, align="r", wrap=False)

    def save(self, path):
        for i, s in enumerate(self.slides, 1):
            if not s.dark and s.max_bottom > 975:
                s.warnings.append(f"content reaches y={s.max_bottom:.0f}px (limit ~930 px, footer at 1000 px)")
        self._footers()
        self.prs.save(str(path))
        return [(i, w) for i, s in enumerate(self.slides, 1) for w in s.warnings]


# --------------------------------------------------------------------------------------
# Shared building blocks
# --------------------------------------------------------------------------------------
def asset(*parts) -> Path:
    return ASSETS.joinpath(*parts)


def logo_dark(s: S, x=1568, y=107, w=280):
    s.pic(asset("logo_dark.png"), x, y, w)


def logo_white(s: S, x=115, y=90, w=286):
    s.pic(asset("logo_white.png"), x, y, w)


def header(s: S, kicker, title, kc=GREY, tw=1300, max_lines=2):
    logo_dark(s)
    s.text(72, 119, 1100, kicker, size=18, weight=600, color=kc, ls=3.5, caps=True, lh=1.2, wrap=False)
    # shrink the title until it fits in max_lines (with 6% safety margin for renderer differences)
    size = 60
    while size > 36 and nlines(title, size, 300, tw * 0.94, -1.2 * size / 60) > max_lines:
        size -= 2
    ls = -1.2 * size / 60
    b = s.text(72, 150, tw, title, size=size, weight=300, color=NAVY, lh=1.1, ls=ls)
    n = nlines(title, size, 300, tw * 0.94, ls)
    if text_w(wrap_lines(title, size, 300, tw, ls)[0], size, 300, ls) > 1480:
        s.warnings.append(f"title runs into the logo - shorten it: '{title[:50]}'")
    if n > max_lines:
        s.warnings.append(f"title has {n} lines (layout allows {max_lines}): '{title[:50]}'")
    return b


def lead(s: S, y, text, w=1500, size=24, color=SLATE, lh=1.45):
    if text:
        return s.text(72, y, w, text, size=size, color=color, lh=lh)
    return y


def kicker15(s, x, y, w, text, color=GREY, size=15, ls=2.5):
    return s.text(x, y, w, text, size=size, weight=700, color=color, ls=ls, caps=True, lh=1.2, wrap=False)


def card(s, x, y, w, h, r=20, fill=PANEL, line=BORDER):
    return s.rect(x, y, w, h, fill=fill, line=line, lw=1, r=r)


def panel_strip(s, y, h, kicker, text, kc=LGREY, size=25, weight=500, ksize=14, color=NAVY, pad_x=36):
    """Key-statement panel (grey, 14px radius)."""
    card(s, 72, y, 1776, h, r=14)
    ktxt = kicker or ""
    kh = ksize * 1.2 if ktxt else 0
    gap = 8 if ktxt else 0
    th_text = th(text, size, weight, 1776 - 2 * pad_x, 1.3)
    top = y + (h - (kh + gap + th_text)) / 2
    if ktxt:
        s.text(72 + pad_x, top, 1000, ktxt, size=ksize, weight=700, color=kc, ls=2.5, caps=True, lh=1.2, wrap=False)
    s.text(72 + pad_x, top + kh + gap, 1776 - 2 * pad_x, text, size=size, weight=weight, color=color, lh=1.3)


def bullets(s, x, y, w, items, size, lh, gap, color=SLATE, marker="—", mcolor=LGREY, indent=40, weight=400):
    for it in items:
        s.text(x, y, indent, marker, size=size, color=mcolor, lh=lh, wrap=False)
        b = s.text(x + indent, y, w - indent, it, size=size, color=color, lh=lh, weight=weight)
        y = b + gap
    return y - gap


def bullets_h(items, w, size, lh, gap, indent, weight=400):
    return sum(th(it, size, weight, w - indent, lh) for it in items) + gap * max(len(items) - 1, 0)


def row_x(n, left=72, total=1776, gap=30):
    cw = (total - gap * (n - 1)) / n
    return cw, [left + i * (cw + gap) for i in range(n)]


def ctext(v, default=""):
    return default if v is None else str(v)


def person(name_or_dict):
    """Resolve a roster name (or dict with overrides) into a full contact record."""
    if isinstance(name_or_dict, str):
        rec = dict(ROSTER.get(name_or_dict, {}))
        rec["name"] = name_or_dict
    else:
        rec = dict(ROSTER.get(name_or_dict.get("name"), {}))
        rec.update(name_or_dict)
    rec.setdefault("role", "")
    rec.setdefault("role_long", rec["role"])
    return rec


def photo_path(rec):
    p = rec.get("photo")
    if not p:
        return None
    pp = Path(p)
    if pp.exists():
        return pp
    cand = asset("team", f"{p}.png")
    return cand if cand.exists() else None


def avatar(s: S, rec, x, y, d):
    pth = photo_path(rec)
    if pth:
        s.pic(pth, x, y, d, d, shape="ellipse")
    else:   # initials fallback
        s.rect(x, y, d, d, fill=BORDER, oval=True)
        ini = "".join(p[0] for p in rec.get("name", "?").split()[:2]).upper()
        s.text(x, y, d, ini, size=d * 0.36, weight=600, color=SLATE, align="c", anchor="m", h=d, lh=1.0, wrap=False)


LAYOUTS: dict = {}


def layout(name):
    def deco(fn):
        LAYOUTS[name] = fn
        return fn
    return deco


def dark_background(deck: Deck) -> S:
    s = deck.new(dark=True)
    s.pic(asset("cover_bg.png"), 0, 0, SLIDE_W, SLIDE_H)
    s.rect(0, 0, SLIDE_W, SLIDE_H, fill=NAVY, fill_alpha=0.35)
    s.gradient(0, 0, SLIDE_W, SLIDE_H, [(0, NAVY, 0.45), (0.5, NAVY, 0.15), (0.75, NAVY, 0.0), (1, NAVY, 0.0)], angle=0)
    return s


# --------------------------------------------------------------------------------------
# F1 Cover
# --------------------------------------------------------------------------------------
@layout("cover")
def cover(deck, d):
    s = dark_background(deck)
    logo_white(s)
    s.text(1805 - 700, 100, 700, d.get("kicker_right", "Family-owned since 1781"), size=18, weight=600, color="C9CEE0",
           ls=3, caps=True, lh=1.2, align="r", wrap=False)
    s.text(115, 296, 1500, d.get("kicker", "Borek Solutions Group"), size=20, weight=700, color=WHITE, ls=4.5, caps=True,
           lh=1.2, wrap=False)
    title = d.get("title", "")
    s.text(115, 346, 1690, title, size=100, weight=300, color=WHITE, lh=1.08, ls=-2)
    if nlines(title, 100, 300, 1690, -2) > 2:
        s.warnings.append("cover headline longer than two lines")
    s.text(115, 584, 930, d.get("intro", ""), size=27, color="E4E7F3", lh=1.4)
    cols = (d.get("columns") or [])[:3]
    for i, c in enumerate(cols):
        x = 115 + i * (310 + 38)
        s.hline(x, 748, 310, WHITE, 1, alpha=0.7)
        b = s.text(x, 748 + 1 + 22, 310, c.get("title", ""), size=24, weight=600, color=WHITE, lh=1.2)
        s.text(x, b + 12, 310, c.get("text", ""), size=20, color="D7DBEA", lh=1.35)
    return s


# --------------------------------------------------------------------------------------
# F2 Who we are (fixed data slide)
# --------------------------------------------------------------------------------------
@layout("who_we_are")
def who_we_are(deck, d):
    s = deck.new()
    logo_dark(s)
    s.text(72, 119, 1100, d.get("kicker", "Who we are · Overview"), size=18, weight=600, color=GREY, ls=3.5, caps=True,
           lh=1.2, wrap=False)
    s.text(72, 150, 1400, d.get("title", "History, locations and leadership"), size=60, weight=300, color=NAVY, lh=1.1,
           ls=-1.2, wrap=False)
    tl = d.get("timeline") or F2_TIMELINE_MASTER
    n = len(tl)
    s.hline(72, 332, 1776, BLUE, 2)
    cw = (1776 - 8 * (n - 1)) / n
    for i, (yr, txt) in enumerate(tl):
        cx = 72 + i * (cw + 8) + cw / 2
        last = i == n - 1
        s.text(cx - cw / 2, 270, cw, yr, size=30, weight=700, color=BLUE if last else NAVY, lh=1.0, align="c",
               wrap=False)
        s.rect(cx - 8, 326, 16, 16, fill=BLUE if last else NAVY, oval=True)
        s.text(cx - 75, 360, 150, txt, size=16, color=SLATE, lh=1.35, align="c")
    kicker15(s, 72, 470, 400, "Locations")
    s.pic(asset("map.png"), 72, 501, 700, alpha=0.85)
    stat = d.get("stat", "250+")
    s.text(72, 681, 700, stat, size=56, weight=700, color=BLUE, ls=-2, lh=1.0, align="c", wrap=False)
    s.text(72, 741, 700, d.get("stat_label", "Employees at four locations"), size=20, weight=600, color=BLUE, lh=1.2,
           align="c", wrap=False)
    locs = d.get("locations") or LOCATIONS
    y = 777

    def loc_runs(name, role):
        return [Run("● " + name, weight=600, color=BLUE), Run(" · " + role, weight=400, color=SLATE)]

    def run_w(name, role):
        return text_w("● " + name, 16, 600) + text_w(" · " + role, 16, 400)
    n0, r0 = locs[0]
    w0 = run_w(n0, r0) + 6
    s.text(72 + 350 - w0 / 2, y, w0 + 20, loc_runs(n0, r0), size=16, lh=1.2, wrap=False)
    rest = locs[1:]
    if rest:
        ws = [run_w(a, b) + 6 for a, b in rest]
        total = sum(ws) + 26 * (len(ws) - 1)
        x = 72 + 350 - total / 2
        for (a, b), wv in zip(rest, ws):
            s.text(x, y + 25, wv + 20, loc_runs(a, b), size=16, lh=1.2, wrap=False)
            x += wv + 26
    kicker15(s, 920, 520, 500, "Management team")
    people = [person(p) for p in (d.get("people") or F2_TEAM)][:9]
    colw = (928 - 48) / 3
    for i, p in enumerate(people):
        cx = 920 + (i % 3) * (colw + 24)
        cy = 567 + (i // 3) * 80
        avatar(s, p, cx, cy, 46)
        s.text(cx + 60, cy + 1, colw - 60, p["name"], size=19, weight=600, color=NAVY, lh=1.2, wrap=False)
        s.text(cx + 60, cy + 25, colw - 60, p["role"], size=15, color=SLATE, lh=1.3, wrap=False)
    s.pic(asset("client_logos.png"), 72, 856, 1776)
    return s


# --------------------------------------------------------------------------------------
# L01 Two-column contrast
# --------------------------------------------------------------------------------------
@layout("contrast")
def contrast(deck, d):
    s = deck.new()
    header(s, d.get("kicker", ""), d.get("title", ""))
    lead(s, 300, d.get("lead"), w=1300, size=26)
    cw = (1776 - 36) / 2
    cards = [d.get("left", {}), d.get("right", {})]
    hs = []
    for c in cards:
        bl = (c.get("bullets") or [])[:4]
        hs.append(36 + 18 + 14 + 41.4 + 30 + bullets_h(bl, cw - 88, 23, 1.4, 16, 40) + 36)
    h = max(360, *hs)
    if 400 + h > 775:
        s.warnings.append("contrast cards too tall; shorten the bullets")
    for i, c in enumerate(cards):
        x = 72 + i * (cw + 36)
        card(s, x, 400, cw, h)
        kicker15(s, x + 44, 436, cw - 88, c.get("label", ""), RED if i == 0 else BLUE)
        b = s.text(x + 44, 436 + 18 + 14, cw - 88, c.get("title", ""), size=36, weight=500, color=NAVY, lh=1.15)
        bullets(s, x + 44, b + 30, cw - 88, (c.get("bullets") or [])[:4], 23, 1.4, 16)
    st = d.get("statement", {})
    panel_strip(s, 785, 110, st.get("kicker", ""), st.get("text", ""))
    if d.get("note"):
        s.text(72, 912, 1776, d["note"], size=19, color=GREY, lh=1.45)
    return s


# --------------------------------------------------------------------------------------
# L02 Four cards + outcome strip
# --------------------------------------------------------------------------------------
@layout("four_cards")
def four_cards(deck, d):
    s = deck.new()
    header(s, d.get("kicker", ""), d.get("title", ""), tw=1628, max_lines=1)
    lead(s, 238, d.get("lead"), w=1545, size=24)
    cards = (d.get("cards") or [])[:4]
    cw, xs = row_x(len(cards), gap=30)
    for c, x in zip(cards, xs):
        card(s, x, 331, cw, 374, r=16)
        b = s.text(x + 34, 365, cw - 68, c.get("title", ""), size=28, weight=600, color=NAVY, lh=1.25)
        s.text(x + 34, b + 16, cw - 68, c.get("text", ""), size=21, color=BODY_ALT, lh=1.5)
    strip = d.get("strip") or {}
    card(s, 72, 737, 1776, 193, r=14)
    kicker15(s, 106, 765, 800, strip.get("kicker", ""), GREY, size=14)
    items = (strip.get("items") or [])[:3]
    iw, ixs = row_x(len(items), 106, 1776 - 68, 30)
    for it, x in zip(items, ixs):
        s.text(x, 765 + 17 + 22, iw, it.get("big", ""), size=44, weight=700, color=NAVY, ls=-1, lh=1.0, wrap=False)
        s.text(x, 765 + 17 + 22 + 44 + 10, iw, it.get("text", ""), size=19, color=SLATE, lh=1.3)
    return s


# --------------------------------------------------------------------------------------
# L03 Comparison matrix
# --------------------------------------------------------------------------------------
@layout("matrix")
def matrix(deck, d):
    s = deck.new()
    header(s, d.get("kicker", ""), d.get("title", ""))
    lead(s, 285, d.get("lead"), w=1607, size=22)
    cols = (d.get("columns") or [])[:5]
    n = max(len(cols), 1)
    cw = (1776 - 600) / n
    y = 362
    hdr_h = 14 * 2 + 14 * 1.3
    last = n - 1
    # navy "our" column background (rounded at the top / bottom)
    rows = d.get("rows") or []
    row_h = 13 * 2 + 21 * 1.3
    total_h = hdr_h + row_h * len(rows)
    if y + total_h > 895:
        s.warnings.append("matrix has too many rows (max ~8)")
    s.rect(72 + 600 + last * cw, y, cw, total_h, fill=NAVY, r=12)
    for i, c in enumerate(cols):
        col = WHITE if i == last else SLATE
        s.text(72 + 600 + i * cw, y + 14, cw, c, size=14, weight=700, color=col, ls=2, caps=True, lh=1.3, align="c")
    s.hline(72, y + hdr_h - 2, 600 + last * cw, NAVY, 2)
    for ri, row in enumerate(rows):
        ry = y + hdr_h + ri * row_h
        s.text(72, ry + 13, 580, row.get("criterion", ""), size=21, color=NAVY, lh=1.3)
        for ci in range(n):
            v = (row.get("cells") or [])[ci] if ci < len(row.get("cells") or []) else ""
            v = ctext(v)
            if v in ("✓", "ok", "yes", "Yes", "✔"):
                v, color, wt = "✓", (WHITE if ci == last else TEAL), 600
            elif v in ("—", "-", "no", "No", ""):
                v, color, wt = "—", (WHITE if ci == last else LGREY), 400
            else:
                color, wt = (WHITE if ci == last else SLATE), (600 if ci == last else 400)
            s.text(72 + 600 + ci * cw, ry + 13, cw, v, size=21, weight=wt, color=color, lh=1.3, align="c")
        if ri < len(rows) - 1:
            s.hline(72, ry + row_h - 1, 600 + last * cw, BORDER, 1)
            s.hline(72 + 600 + last * cw, ry + row_h - 1, cw, WHITE, 1, alpha=0.15)
    if d.get("statement"):
        s.text(72, 905, 1500, d["statement"], size=24, weight=600, color=NAVY, lh=1.35)
    return s


# --------------------------------------------------------------------------------------
# L04 Three offers
# --------------------------------------------------------------------------------------
@layout("offers")
def offers(deck, d):
    s = deck.new()
    header(s, d.get("kicker", ""), d.get("title", ""), max_lines=1)
    lead(s, 250, d.get("lead"), w=1500, size=24)
    cards = (d.get("cards") or [])[:3]
    cw, xs = row_x(len(cards), gap=36)
    iw = cw - 72
    ends = []                               # cards grow with their content (min-height 420), as in the HTML grid
    for c in cards:
        t_end = 369 + 18 + 22 + th(c.get("title", ""), 38, 500, iw, 1.15) + 16 + th(c.get("text", ""), 22, 400, iw, 1.45)
        ends.append(t_end + 26 + bullets_h((c.get("bullets") or [])[:3], iw, 21, 1.3, 12, 20))
    rule_y = max(max(ends + [0]), 335 + 420 - 34 - 41 - 24)
    h = rule_y + 1 + 24 + 41 + 34 - 335
    for i, (c, x) in enumerate(zip(cards, xs)):
        col = c.get("color") or ACCENTS[i % 3]
        card(s, x, 335, cw, h)
        kicker15(s, x + 36, 369, iw, c.get("kicker", f"0{i + 1}"), col)
        b = s.text(x + 36, 369 + 18 + 22, iw, c.get("title", ""), size=38, weight=500, color=NAVY, lh=1.15)
        b = s.text(x + 36, b + 16, iw, c.get("text", ""), size=22, color=SLATE, lh=1.45)
        bullets((s), x + 36, b + 26, iw, (c.get("bullets") or [])[:3], 21, 1.3, 12, marker="·", mcolor=col, indent=20)
        s.hline(x + 36, rule_y, iw, BORDER, 1)
        s.text(x + 36, rule_y + 25, iw, c.get("price", ""), size=34, weight=300, color=NAVY, ls=-1, lh=1.2, wrap=False)
    p = d.get("panel") or {}
    panel_strip(s, 775, 110, p.get("kicker", ""), p.get("text", ""))
    if d.get("note"):
        s.text(72, 903, 1776, d["note"], size=19, color=GREY, lh=1.45)
    return s


# --------------------------------------------------------------------------------------
# L05 Five pillars
# --------------------------------------------------------------------------------------
@layout("pillars")
def pillars(deck, d):
    s = deck.new()
    header(s, d.get("kicker", ""), d.get("title", ""))
    lead(s, 300, d.get("lead"), w=1400, size=24)
    cards = (d.get("cards") or [])[:5]
    cw, xs = row_x(len(cards), gap=26)
    pal = [BLUE, TEAL, RED, ORANGE, BLUE]
    for i, (c, x) in enumerate(zip(cards, xs)):
        card(s, x, 395, cw, 395, r=16)
        kicker15(s, x + 30, 425, cw - 60, c.get("n", f"{i + 1:02d}"), pal[i % 5])
        b = s.text(x + 30, 425 + 18 + 16, cw - 60, c.get("title", ""), size=27, weight=600, color=NAVY, lh=1.25)
        s.text(x + 30, b + 14, cw - 60, c.get("text", ""), size=19, color=SLATE, lh=1.45)
    st = d.get("statement") or {}
    panel_strip(s, 813, 117, st.get("kicker", ""), st.get("text", ""))
    return s


# --------------------------------------------------------------------------------------
# L06 Section divider
# --------------------------------------------------------------------------------------
@layout("divider")
def divider(deck, d):
    s = dark_background(deck)
    logo_white(s)
    for i, c in enumerate(ACCENTS):
        s.rect(115 + i * 153, 360, 153, 7, fill=c)
    s.text(115, 300, 1500, d.get("kicker", ""), size=20, weight=700, color="C9CEE0", ls=4.5, caps=True, lh=1.2,
           wrap=False)
    if nlines(d.get("title", ""), 105, 300, 1690, -2.5) > 1:
        s.warnings.append("divider title should fit on one line")
    s.text(115, 420, 1690, d.get("title", ""), size=105, weight=300, color=WHITE, lh=1.05, ls=-2.5)
    s.text(115, 585, 1430, d.get("text", ""), size=30, color="E4E7F3", lh=1.4)
    return s


# --------------------------------------------------------------------------------------
# L07 Three-stage process
# --------------------------------------------------------------------------------------
@layout("process")
def process(deck, d):
    s = deck.new()
    header(s, d.get("kicker", ""), d.get("title", ""), max_lines=1)
    lead(s, 250, d.get("lead"), w=1500, size=24)
    st = (d.get("stages") or [])[:3]
    cw, xs = row_x(len(st), gap=36)
    iw = cw - 80
    pal = [BLUE, TEAL, RED]

    def res_txt(c):
        return ((c.get("result_label", "") + " ") if c.get("result_label") else "") + c.get("result", "")
    p_end = [366 + 46 + 36 + th(c.get("title", ""), 39, 600, iw, 1.1) + 22 + th(c.get("text", ""), 21, 400, iw, 1.45)
             for c in st]
    rh = max([th(res_txt(c), 20, 400, iw, 1.4) for c in st] + [28])
    h = max(410, max(p_end + [0]) + 12 + 1 + 22 + rh + 36 - 330)     # cards grow with their content
    for i, (c, x) in enumerate(zip(st, xs)):
        card(s, x, 330, cw, h)
        s.text(x + 40, 366, iw, c.get("n", f"{i + 1:02d}"), size=46, weight=600, color=pal[i % 3], lh=1.0, wrap=False)
        b = s.text(x + 40, 366 + 46 + 36, iw, c.get("title", ""), size=39, weight=600, color=NAVY, lh=1.1)
        s.text(x + 40, b + 22, iw, c.get("text", ""), size=21, color=SLATE, lh=1.45)
        ry = 330 + h - 36 - rh
        s.hline(x + 40, ry - 23, iw, RULE, 1)
        s.text(x + 40, ry, iw, [Run((c.get("result_label", "") + " ") if c.get("result_label") else "", weight=600,
                                   color=NAVY), Run(c.get("result", ""), color=NAVY)], size=20, color=NAVY, lh=1.4)
    p = d.get("panel") or {}
    panel_strip(s, 765, 110, p.get("kicker", ""), p.get("text", ""))
    if d.get("note"):
        s.text(72, 903, 1776, d["note"], size=19, color=GREY, lh=1.45)
    return s


# --------------------------------------------------------------------------------------
# L08 Matrix + notes
# --------------------------------------------------------------------------------------
@layout("matrix_notes")
def matrix_notes(deck, d):
    s = deck.new()
    header(s, d.get("kicker", ""), d.get("title", ""), kc=BLUE)
    lead(s, 300, d.get("lead"), w=1300, size=24)
    rows = (d.get("rows") or [])[:7]
    y = 400
    for i, r in enumerate(rows):
        a = r.get("code", "")
        name = r.get("title", "")
        rh = 20 * 2 + 22 * 1.3
        runs = [Run(a + "   ", weight=600, color=BLUE), Run(name, weight=600, color=NAVY)] if a else [Run(name, weight=600, color=NAVY)]
        s.text(72, y + 20, 340, runs, size=22, lh=1.3)
        s.text(72 + 360, y + 20, 720, r.get("text", ""), size=22, color=SLATE, lh=1.3)
        if i < len(rows) - 1:
            s.hline(72, y + rh - 1, 1080, BORDER, 1)
        y += rh
    ny = 400
    for n in (d.get("notes") or [])[:2]:
        h = 32 * 2 + th(n.get("title", ""), 28, 600, 536, 1.25) + 14 + th(n.get("text", ""), 20, 400, 536, 1.45)
        card(s, 1240, ny, 608, h)
        b = s.text(1240 + 36, ny + 32, 536, n.get("title", ""), size=28, weight=600, color=NAVY, lh=1.25)
        s.text(1240 + 36, b + 14, 536, n.get("text", ""), size=20, color=SLATE, lh=1.45)
        ny += h + 28
    return s


# --------------------------------------------------------------------------------------
# L10 Data table + points
# --------------------------------------------------------------------------------------
@layout("data_table")
def data_table(deck, d):
    s = deck.new()
    header(s, d.get("kicker", ""), d.get("title", ""), kc=BLUE)
    lead(s, 300, d.get("lead"), w=1400, size=23)
    kicker15(s, 72, 395, 900, d.get("table_label", ""))
    heads = d.get("headers") or []
    rows = d.get("rows") or []
    ncol = max(len(heads), 1)
    c0 = 1100 / (1.4 + (ncol - 1))
    widths = [c0 * 1.4] + [c0] * (ncol - 1)
    xs = [72 + sum(widths[:i]) for i in range(ncol)]
    y = 430
    for i, hd in enumerate(heads):
        s.text(xs[i], y + 12, widths[i], hd, size=14, weight=700, color=SLATE, ls=2, caps=True, lh=1.3,
               align="l" if i == 0 else "r")
    y += 12 * 2 + 14 * 1.3
    s.hline(72, y - 2, 1100, NAVY, 2)
    for ri, row in enumerate(rows):
        total = ri == len(rows) - 1 and len(rows) > 1
        pad = 14 if total else 13
        for i in range(ncol):
            v = ctext(row[i]) if i < len(row) else ""
            if total:
                wt, col = 700, (BLUE if i == ncol - 1 else (NAVY if i < ncol - 1 and i != ncol - 2 else SLATE))
                col = BLUE if i == ncol - 1 else NAVY
            else:
                wt = 600 if i in (0, ncol - 1) else 400
                col = SLATE if i == ncol - 2 and ncol > 3 else NAVY
            s.text(xs[i], y + pad, widths[i], v, size=22, weight=wt, color=col, lh=1.3, align="l" if i == 0 else "r")
        rh = pad * 2 + 22 * 1.3
        y += rh
        if ri == len(rows) - 2:
            s.hline(72, y - 2, 1100, NAVY, 2)
        elif ri < len(rows) - 1:
            s.hline(72, y - 1, 1100, BORDER, 1)
    py = 430
    for p in (d.get("points") or [])[:3]:
        b = s.text(1260, py + 22, 588, p.get("title", ""), size=26, weight=600, color=NAVY, lh=1.25)
        b = s.text(1260, b + 10, 588, p.get("text", ""), size=20, color=SLATE, lh=1.45)
        py = b + 30
    if d.get("callout"):
        h = 44 + th(d["callout"], 21, 600, 536, 1.35)
        card(s, 1260, py, 588, h, r=14)
        s.text(1260 + 26, py + 22, 536, d["callout"], size=21, weight=600, color=NAVY, lh=1.35)
    return s


# --------------------------------------------------------------------------------------
# L11 Big number + bars
# --------------------------------------------------------------------------------------
@layout("big_number")
def big_number(deck, d):
    s = deck.new()
    header(s, d.get("kicker", ""), d.get("title", ""), kc=BLUE)
    lead(s, 300, d.get("lead"), w=1400, size=23)
    bn = d.get("big") or {}
    kicker15(s, 72, 420, 380, bn.get("label", ""))
    bsz = 180
    while bsz > 60 and text_w(str(bn.get("value", "")), bsz, 700, -8 * bsz / 180) > 470:
        bsz -= 4                                     # long numbers shrink so they never reach the bars
    s.text(72, 459 + (180 - bsz) * 0.8, 480, bn.get("value", ""), size=bsz, weight=700, color=NAVY,
           ls=-8 * bsz / 180, lh=1.0, wrap=False)
    s.text(72, 459 + 180 + 14, 380, bn.get("text", ""), size=22, color=SLATE, lh=1.4)
    kicker15(s, 560, 420, 700, d.get("bars_label", ""))
    y = 420 + 18 + 34
    bw = 1288 - 280 - 170 - 48
    pal = [BLUE, TEAL, RED]
    for i, b in enumerate((d.get("bars") or [])[:4]):
        tsz = 28
        while tsz > 18 and text_w(b.get("title", ""), tsz, 600) > 280 * 0.94:
            tsz -= 1                                 # keep bar titles on one line so they never hit the text below
        s.text(560, y, 280, b.get("title", ""), size=tsz, weight=600, color=NAVY, lh=1.2)
        desc_bottom = s.text(560, y + 34 + 4, 280, b.get("text", ""), size=19, color=SLATE, lh=1.3)
        bx = 560 + 280 + 24
        by = y + (61 - 56) / 2
        s.rect(bx, by, bw, 56, fill="F1F3F8", r=10)
        pct = max(0.0, min(100.0, float(b.get("pct", 100))))
        s.rect(bx, by, max(bw * pct / 100, 12), 56, fill=b.get("color") or pal[i % 3], r=10)
        s.text(bx + bw + 24, y + 12, 170, b.get("value", ""), size=30, weight=600, color=NAVY, lh=1.2, wrap=False)
        y += max(61 + 34, desc_bottom - y + 24)      # grow the row if the description wraps to many lines
    if d.get("note"):
        s.text(72, 870, 1600, d["note"], size=23, color=SLATE, lh=1.4)
    return s


# --------------------------------------------------------------------------------------
# L12 Flow + facts
# --------------------------------------------------------------------------------------
@layout("flow_facts")
def flow_facts(deck, d):
    s = deck.new()
    header(s, d.get("kicker", ""), d.get("title", ""), kc=BLUE)
    lead(s, 300, d.get("lead"), w=1500, size=23)
    steps = (d.get("steps") or [])[:5]
    cw, xs = row_x(len(steps), gap=20)
    for i, (t, x) in enumerate(zip(steps, xs)):
        lines = nlines(t, 26, 600, cw - 60)
        h = 60 + lines * 26 * 1.3
        y = 430 + (92 - h) / 2 if h < 92 else 430
        if i == 0:
            s.rect(x, y, cw, h, fill=NAVY, r=16)
        else:
            card(s, x, y, cw, h, r=16)
        s.text(x + 30, y + 30, cw - 60, t, size=26, weight=600, color=WHITE if i == 0 else NAVY, lh=1.3, align="c")
    facts = (d.get("facts") or [])[:4]
    fw, fxs = row_x(len(facts), gap=30)
    for f, x in zip(facts, fxs):
        s.text(x, 630, fw, f.get("big", ""), size=54, weight=700, color=NAVY, ls=-2, lh=1.0, wrap=False)
        s.text(x, 630 + 54 + 14, fw, f.get("text", ""), size=20, color=SLATE, lh=1.45)
    if d.get("note"):
        s.text(72, 860, 1600, d["note"], size=23, color=SLATE, lh=1.4)
    return s


# --------------------------------------------------------------------------------------
# L14 Objections / FAQ
# --------------------------------------------------------------------------------------
@layout("faq")
def faq(deck, d):
    s = deck.new()
    header(s, d.get("kicker", ""), d.get("title", ""), kc=BLUE, max_lines=1)
    lead(s, 250, d.get("lead"), w=1500, size=23)
    items = (d.get("items") or [])[:4]
    cw = (1776 - 36) / 2
    y = 390
    for r in range(0, len(items), 2):
        pair = items[r:r + 2]
        qs = []
        for it in pair:
            q = it.get("q", "")
            if not q.startswith(("“", '"', "„")):
                q = f"“{q.strip('“”')}”"
            qs.append(q)
        hs = [34 * 2 + th(q, 32, 300, cw - 80, 1.25) + 16 + th(it.get("a", ""), 21, 400, cw - 80, 1.45)
              for q, it in zip(qs, pair)]
        h = max(hs)
        for i, (q, it) in enumerate(zip(qs, pair)):
            x = 72 + i * (cw + 36)
            card(s, x, y, cw, h)
            b = s.text(x + 40, y + 34, cw - 80, q, size=32, weight=300, color=NAVY, lh=1.25, italic=True)
            s.text(x + 40, b + 16, cw - 80, it.get("a", ""), size=21, color=SLATE, lh=1.45)
        y += h + 30
    if y - 30 > 860:
        s.warnings.append("FAQ cards too tall; shorten the answers")
    if d.get("support"):
        s.text(72, 870, 1600, d["support"], size=23, weight=600, color=NAVY, lh=1.4)
    return s


# --------------------------------------------------------------------------------------
# L15 Two models + features
# --------------------------------------------------------------------------------------
@layout("two_models")
def two_models(deck, d):
    s = deck.new()
    header(s, d.get("kicker", ""), d.get("title", ""), kc=RED)
    lead(s, 300, d.get("lead"), w=1500, size=23)
    y = 410
    for i, m in enumerate((d.get("models") or [])[:2]):
        h = 30 * 2 + 14 * 1.2 + 10 + th(m.get("title", ""), 30, 600, 688, 1.2) + 12 + th(m.get("text", ""), 20, 400, 688, 1.45)
        card(s, 72, y, 760, h)
        kicker15(s, 108, y + 30, 688, m.get("kicker", ""), BLUE if i == 0 else RED, size=14, ls=2)
        b = s.text(108, y + 30 + 17 + 10, 688, m.get("title", ""), size=30, weight=600, color=NAVY, lh=1.2)
        s.text(108, b + 12, 688, m.get("text", ""), size=20, color=SLATE, lh=1.45)
        y += h + 24
    feats = (d.get("features") or [])[:4]
    colw = (928 - 40) / 2
    yy = 410
    for r in range(0, len(feats), 2):
        row = feats[r:r + 2]
        hmax = 0
        for i, f in enumerate(row):
            x = 920 + i * (colw + 40)
            s.hline(x, yy, colw, RULE, 1)
            b = s.text(x, yy + 22, colw, f.get("title", ""), size=25, weight=600, color=NAVY, lh=1.25)
            b = s.text(x, b + 10, colw, f.get("text", ""), size=19, color=SLATE, lh=1.45)
            hmax = max(hmax, b - yy)
        yy += hmax + 30
    if d.get("note"):
        s.text(72, 903, 1776, d["note"], size=19, color=GREY, lh=1.45)
    return s


# --------------------------------------------------------------------------------------
# L16 Phased steps
# --------------------------------------------------------------------------------------
@layout("phased")
def phased(deck, d):
    s = deck.new()
    header(s, d.get("kicker", ""), d.get("title", ""), max_lines=1)
    lead(s, 250, d.get("lead"), w=1500, size=23)
    steps = (d.get("steps") or [])[:4]
    cw, xs = row_x(len(steps), gap=30)
    for i, (c, x) in enumerate(zip(steps, xs)):
        card(s, x, 380, cw, 395, r=16)
        kicker15(s, x + 32, 410, cw - 64, c.get("when", ""), ACCENTS[i % 4])
        b = s.text(x + 32, 410 + 18 + 14, cw - 64, c.get("title", ""), size=30, weight=600, color=NAVY, lh=1.2)
        s.text(x + 32, b + 16, cw - 64, c.get("text", ""), size=20, color=SLATE, lh=1.45)
    p = d.get("panel") or {}
    h = 28 * 2 + 18 + 14 + th(p.get("text", ""), 24, 400, 1704, 1.4)
    card(s, 72, 805, 1776, h, r=14)
    kicker15(s, 108, 833, 800, p.get("kicker", ""))
    s.text(108, 833 + 18 + 14, 1704, p.get("text", ""), size=24, color=NAVY, lh=1.4)
    return s


# --------------------------------------------------------------------------------------
# L17 Three packages
# --------------------------------------------------------------------------------------
@layout("packages")
def packages(deck, d):
    s = deck.new()
    header(s, d.get("kicker", ""), d.get("title", ""), kc=RED)
    lead(s, 250, d.get("lead"), w=1500, size=23)
    pk = (d.get("packages") or [])[:3]
    cw, xs = row_x(len(pk), gap=36)
    hi = d.get("highlight", 1)
    for i, (c, x) in enumerate(zip(pk, xs)):
        on = i == hi
        fg, sub, bul = (WHITE, "C9CEE0", "DDE1EE") if on else (NAVY, SLATE, SLATE)
        if on:
            s.rect(x, 335, cw, 435, fill=NAVY, line=NAVY, lw=1, r=20)
        else:
            card(s, x, 335, cw, 435)
        iw = cw - 72
        tw = iw - (150 if on else 0)
        b = s.text(x + 36, 369, tw, c.get("title", ""), size=32, weight=500, color=fg, lh=1.2)
        b = s.text(x + 36, b + 8, iw, c.get("sub", ""), size=20, color=sub, lh=1.2)
        b = s.text(x + 36, b + 26, iw, c.get("price", ""), size=66, weight=300, color=fg, ls=-2, lh=1.0, wrap=False)
        b = s.text(x + 36, b + 8, iw, c.get("unit", ""), size=20, color=sub, lh=1.2)
        s.hline(x + 36, b + 20, iw, WHITE if on else BORDER, 1, alpha=0.22 if on else None)
        bullets(s, x + 36, b + 20 + 1 + 18, iw, (c.get("bullets") or [])[:4], 22, 1.3, 12, color=bul, marker="·",
                mcolor=bul, indent=28)
        if on:
            bd = d.get("badge", "Popular")
            bwid = text_w(bd.upper(), 13, 700, 2) + 36
            s.rect(x + cw - 34 - bwid, 369, bwid, 36, fill=WHITE, r=999)
            s.text(x + cw - 34 - bwid, 369, bwid, bd, size=13, weight=700, color=NAVY, ls=2, caps=True, lh=1.0,
                   align="c", anchor="m", h=36, wrap=False)
    p = d.get("panel") or {}
    card(s, 72, 790, 1776, 140, r=14)
    kh = 14 * 1.2 + 8 + th(p.get("text", ""), 23, 400, 1704, 1.4)
    top = 790 + (140 - kh) / 2
    kicker15(s, 108, top, 800, p.get("kicker", ""), BLUE, size=14)
    s.text(108, top + 14 * 1.2 + 8, 1704, p.get("text", ""), size=23, color=NAVY, lh=1.4)
    return s


# --------------------------------------------------------------------------------------
# L18 Capabilities + stats
# --------------------------------------------------------------------------------------
@layout("capabilities")
def capabilities(deck, d):
    s = deck.new()
    header(s, d.get("kicker", ""), d.get("title", ""))
    lead(s, 300, d.get("lead"), w=1500, size=23)
    cards = (d.get("cards") or [])[:3]
    cw, xs = row_x(len(cards), gap=30)
    h = max([32 * 2 + th(c.get("title", ""), 28, 600, cw - 72, 1.2) + 14 + th(c.get("text", ""), 20, 400, cw - 72, 1.45)
             for c in cards] + [200])
    for c, x in zip(cards, xs):
        card(s, x, 420, cw, h, fill="F7F8FC")
        b = s.text(x + 36, 452, cw - 72, c.get("title", ""), size=28, weight=600, color=NAVY, lh=1.2)
        s.text(x + 36, b + 14, cw - 72, c.get("text", ""), size=20, color=SLATE, lh=1.45)
    stats = (d.get("stats") or [])[:5]
    sw, sxs = row_x(len(stats), gap=30)
    for t, x in zip(stats, sxs):
        s.text(x, 724, sw, t.get("big", ""), size=56, weight=700, color=NAVY, ls=-2, lh=1.0, wrap=False)
        s.text(x, 724 + 56 + 12, sw, t.get("text", ""), size=18, color=SLATE, lh=1.4)
    if d.get("note"):
        s.text(72, 903, 1776, d["note"], size=19, color=GREY, lh=1.45)
    return s


# --------------------------------------------------------------------------------------
# L19 Certificates + commitments
# --------------------------------------------------------------------------------------
@layout("certificates")
def certificates(deck, d):
    s = deck.new()
    header(s, d.get("kicker", ""), d.get("title", ""))
    lead(s, 300, d.get("lead"), w=1500, size=23)
    bd = (d.get("badges") or [])[:4]
    cw, xs = row_x(len(bd), gap=24)
    for b, x in zip(bd, xs):
        h = 26 * 2 + 34 * 1.2 + 8 + th(b.get("text", ""), 18, 400, cw - 60, 1.3)
        s.rect(x, 420, cw, max(h, 126), fill=NAVY, r=16)
        s.text(x + 30, 446, cw - 60, b.get("title", ""), size=34, weight=700, color=WHITE, ls=-1, lh=1.2, wrap=False)
        s.text(x + 30, 446 + 41 + 8, cw - 60, b.get("text", ""), size=18, color="C9CEE0", lh=1.3)
    cm = (d.get("commitments") or [])[:6]
    cw3 = (1776 - 80) / 3
    y = 600
    for r in range(0, len(cm), 3):
        row = cm[r:r + 3]
        hmax = 0
        for i, c in enumerate(row):
            x = 72 + i * (cw3 + 40)
            s.hline(x, y, cw3, RULE, 1)
            b = s.text(x, y + 22, cw3, c.get("title", ""), size=25, weight=600, color=NAVY, lh=1.25)
            b = s.text(x, b + 10, cw3, c.get("text", ""), size=19, color=SLATE, lh=1.45)
            hmax = max(hmax, b - y)
        y += hmax + 34
    return s


# --------------------------------------------------------------------------------------
# L20 Option comparison
# --------------------------------------------------------------------------------------
@layout("options")
def options(deck, d):
    s = deck.new()
    header(s, d.get("kicker", ""), d.get("title", ""), kc=TEAL)
    opts = (d.get("options") or [])[:3]
    n = len(opts)
    colw = (1776 - 260 - 24 * n) / max(n, 1)                       # grid: 260 px label column + n columns, 24 px gaps
    xs = [72 + 260 + 24 + i * (colw + 24) for i in range(n)]
    y = 300
    pal = [BLUE, TEAL, RED]
    names = [o if isinstance(o, str) else o.get("name", "") for o in opts]
    head_h = max([16 * 2 + th(t, 22, 600, colw, 1.35) for t in names] + [60])
    s.hline(72, y + head_h - 2, 260, NAVY, 2)
    for i, name in enumerate(names):
        s.text(xs[i], y + 16, colw, name, size=22, weight=600, color=pal[i % 3], lh=1.35)
        s.hline(xs[i], y + head_h - 4, colw, pal[i % 3], 4)
    y += head_h
    for row in d.get("rows") or []:
        cells = row.get("cells") or []
        style = row.get("style")
        wt, col = (600, NAVY) if style == "bold" else (400, SLATE if style == "muted" else NAVY)
        rh = 16 * 2 + max([th(c, 21, wt, colw, 1.35) for c in cells] + [28.35])
        s.text(72, y + 16, 250, row.get("label", ""), size=14, weight=700, color=SLATE, ls=2, caps=True, lh=1.35)
        s.hline(72, y + rh - 1, 260, BORDER, 1)
        for i in range(n):
            s.text(xs[i], y + 16, colw, ctext(cells[i]) if i < len(cells) else "", size=21, weight=wt, color=col, lh=1.35)
            s.hline(xs[i], y + rh - 1, colw, BORDER, 1)
        y += rh
    av = d.get("availability")
    if av:
        s.text(72, y + 16, 250, d.get("availability_label", "Availability"), size=14, weight=700, color=SLATE, ls=2,
               caps=True, lh=1.35)
        for i in range(n):
            a = av[i] if i < len(av) else {}
            if isinstance(a, str):
                a = {"text": a, "on": True}
            t = a.get("text", "")
            pw = text_w(t.upper(), 14, 700, 1.5) + 28
            on = a.get("on", True)
            s.rect(xs[i], y + 16, pw, 14 * 1.2 + 12, fill=NAVY if on else "EEF2F7", r=999)
            s.text(xs[i], y + 16, pw, t, size=14, weight=700, color=WHITE if on else SLATE, ls=1.5, caps=True, lh=1.0,
                   align="c", anchor="m", h=14 * 1.2 + 12, wrap=False)
        y += 16 * 2 + 14 * 1.2 + 12
    if y > 880:
        s.warnings.append("options table too tall; shorten the cells")
    if d.get("recommendation"):
        s.text(72, 900, 1600, d["recommendation"], size=23, weight=600, color=NAVY, lh=1.4)
    return s


# --------------------------------------------------------------------------------------
# L21 Price tables
# --------------------------------------------------------------------------------------
def _table(s, x, y, widths, gap, heads, rows, size, bold_cols=(), right_cols=(), muted_cols=()):
    xs = [x]
    for w in widths[:-1]:
        xs.append(xs[-1] + w + gap)
    hh = 10 * 2 + 13 * 1.35
    for i, hd in enumerate(heads):
        s.text(xs[i], y + 10, widths[i], hd, size=13, weight=700, color=SLATE, ls=2, caps=True, lh=1.35,
               align="r" if i in right_cols else "l")
    s.hline(x, y + hh - 2, sum(widths) + gap * (len(widths) - 1), NAVY, 2)
    y += hh
    for ri, row in enumerate(rows):
        hmax = max(th(ctext(row[i]) if i < len(row) else "", size, 600 if i in bold_cols else 400, widths[i], 1.35)
                   for i in range(len(widths)))
        for i in range(len(widths)):
            v = ctext(row[i]) if i < len(row) else ""
            s.text(xs[i], y + 14, widths[i], v, size=size, weight=600 if i in bold_cols else 400,
                   color=NAVY if i in bold_cols else SLATE, lh=1.35, align="r" if i in right_cols else "l")
        y += 14 * 2 + hmax
        if ri < len(rows) - 1:
            s.hline(x, y - 1, sum(widths) + gap * (len(widths) - 1), BORDER, 1)
    return y


@layout("price_tables")
def price_tables(deck, d):
    s = deck.new()
    logo_dark(s)
    s.text(72, 119, 1100, d.get("kicker", ""), size=18, weight=600, color=GREY, ls=3.5, caps=True, lh=1.2, wrap=False)
    s.text(72, 150, 1560, d.get("title", ""), size=60, weight=300, color=NAVY, lh=1.1, ls=-1.2, wrap=False)
    lead(s, 250, d.get("lead"), w=1500, size=21)
    L, R = d.get("left") or {}, d.get("right") or {}
    kicker15(s, 72, 320, 800, L.get("kicker", ""), RED)
    _table(s, 72, 320 + 18 + 18, [230, 800 - 230 - 150 - 40, 150], 20, L.get("headers", []), L.get("rows", []), 19,
           bold_cols=(0, 2), right_cols=(2,))
    kicker15(s, 940, 320, 800, R.get("kicker", ""), TEAL)
    rw = [170, 150, 130, 180]
    rw.append(908 - sum(rw) - 16 * 4)
    _table(s, 940, 320 + 18 + 18, rw, 16, R.get("headers", []), R.get("rows", []), 18, bold_cols=(0, 3))
    if d.get("note"):
        s.text(72, 920, 1776, d["note"], size=19, color=GREY, lh=1.45)
    return s


# --------------------------------------------------------------------------------------
# L22 Case study
# --------------------------------------------------------------------------------------
@layout("case_study")
def case_study(deck, d):
    s = deck.new()
    header(s, d.get("kicker", ""), d.get("title", ""))
    img = d.get("image")
    if img and Path(img).exists():
        s.pic(img, 72, 300, 820, 600, shape="round", radius=14)
    else:
        s.rect(72, 300, 820, 600, fill=PANEL, r=14)
        s.text(72, 300, 820, "Drop a product screenshot (820 × 600)", size=20, color=GREY, align="c", anchor="m", h=600,
               lh=1.2)
    kicker15(s, 960, 300, 700, d.get("problem_label", "Problem → solution"))
    b = s.text(960, 300 + 18 + 26, 888, d.get("text", ""), size=26, color=SLATE, lh=1.45)
    res = (d.get("results") or [])[:2]
    y = b + 26 + 6
    h = max([30 * 2 + 15 * 1.2 + 14 + 63 + 10 + th(r.get("text", ""), 21, 400, 364, 1.3) for r in res] + [0])
    for i, r in enumerate(res):
        x = 960 + i * (432 + 24)
        card(s, x, y, 432, h, r=14)
        kicker15(s, x + 34, y + 30, 364, r.get("label", "Result"), GREY, size=15, ls=2)
        s.text(x + 34, y + 30 + 18 + 14, 364, r.get("value", ""), size=63, weight=700, color=NAVY, ls=-2, lh=1.0, wrap=False)
        s.text(x + 34, y + 30 + 18 + 14 + 63 + 10, 364, r.get("text", ""), size=21, color=SLATE, lh=1.3)
    return s


# --------------------------------------------------------------------------------------
# L23 Management team
# --------------------------------------------------------------------------------------
@layout("team")
def team(deck, d):
    s = deck.new()
    header(s, d.get("kicker", "Management team"), d.get("title", "The people who stand behind every delivery"))
    people = [person(p) for p in (d.get("people") or F2_TEAM + ["Max Fichtner"])][:8]
    cw = (1776 - 3 * 36) / 4
    y = 340
    for r in range(0, len(people), 4):
        row = people[r:r + 4]
        for i, p in enumerate(row):
            x = 72 + i * (cw + 36)
            s.hline(x, y, cw, BORDER, 1)
            avatar(s, p, x, y + 27, 130)
            b = s.text(x, y + 27 + 130 + 22, cw, p["name"], size=28, weight=600, color=NAVY, lh=1.2)
            s.text(x, b + 6, cw, p["role"], size=19, color=SLATE, lh=1.3)
        y += 27 + 130 + 22 + 34 + 6 + 25 + 40
    return s


# --------------------------------------------------------------------------------------
# L24 Locations map
# --------------------------------------------------------------------------------------
@layout("locations")
def locations(deck, d):
    s = deck.new()
    header(s, d.get("kicker", "Locations"), d.get("title", "Four locations, one delivery standard"))
    s.pic(asset("map.png"), 310, 250, 1300)
    s.text(310, 520, 1300, d.get("stat", "250+"), size=170, weight=700, color=BLUE, ls=-6, lh=1.0, align="c", wrap=False)
    s.text(310, 520 + 170 + 14, 1300, d.get("stat_label", "Employees at four locations"), size=44, weight=600, color=BLUE,
           ls=-0.5, lh=1.2, align="c", wrap=False)
    locs = d.get("locations") or LOCATIONS

    def item_w(a, b):
        return 16 + 14 + text_w(a, 36, 700) + 14 + text_w("· " + b, 30, 400) + 8

    def draw(a, b, x, y):
        s.rect(x, y + 14, 16, 16, fill=NAVY, oval=True)
        s.text(x + 30, y, item_w(a, b), [Run(a, size=36, weight=700, color=NAVY), Run("  · " + b, size=30, weight=400, color="6B7285")],
               size=36, lh=1.2, wrap=False)
    a0, b0 = locs[0]
    w0 = item_w(a0, b0)
    draw(a0, b0, 960 - w0 / 2, 790)
    rest = locs[1:]
    if rest:
        ws = [item_w(a, b) for a, b in rest]
        tot = sum(ws) + 64 * (len(ws) - 1)
        x = 960 - tot / 2
        for (a, b), w_ in zip(rest, ws):
            draw(a, b, x, 790 + 43 + 22)
            x += w_ + 64
    return s


# --------------------------------------------------------------------------------------
# F3 Closing
# --------------------------------------------------------------------------------------
@layout("closing")
def closing(deck, d):
    s = dark_background(deck)
    s.footer = "closing"
    logo_white(s, 90, 130, 314)
    s.text(90, 272, 1100, d.get("title", "Let’s talk"), size=64, weight=300, color=WHITE, ls=-1.2, lh=1.1)
    cons = [person(c) for c in (d.get("contacts") or ["Konstantin Borek", "Viktoria Schünemann", "Fiona Oldenburg",
                                                       "Max Fichtner"])][:4]
    n = len(cons)
    cw = (1740 - 30 * (n - 1)) / n if n else 412.5
    cw = min(cw, 412.5)
    long_roles = [nlines(c["role_long"], 19, 600, cw - 50) for c in cons]
    h = max(281, 17 + 76 + 18 + 28.8 + 4 + max(long_roles + [1]) * 19 * 1.3 + 16 + 50 + 37)
    for i, c in enumerate(cons):
        x = 90 + i * (cw + 30)
        s.rect(x, 432, cw, h, fill=WHITE, fill_alpha=0.06, line=WHITE, line_alpha=0.14, lw=1.6, r=18.5)
        avatar(s, c, x + 25, 432 + 17, 76)
        b = s.text(x + 25, 432 + 17 + 76 + 18, cw - 50, c["name"], size=24, weight=700, color=WHITE, lh=1.2)
        b = s.text(x + 25, b + 4, cw - 50, c["role_long"] if d.get("long_roles", True) else c["role"], size=19,
                   weight=600, color="B3B3B3", lh=1.3)
        yy = b + 16
        for icon, key, sz in (("mail", "email", 15), ("phone", "phone", 17)):
            if c.get(key):
                s.pic(asset("icons", f"{icon}.png"), x + 25, yy, 20, 20)
                s.text(x + 25 + 32, yy + 1, cw - 50 - 32, c[key], size=sz, color="DEDEDE", lh=1.2, wrap=False)
                yy += 30
    s.text(90, 782 if h <= 300 else 432 + h + 40, 1500, d.get("tagline", ""), size=24, weight=300, color=WHITE, lh=1.45,
           wrap=False)
    y = 906
    for icon, txt in (("pin", d.get("address", ADDRESS)), ("phone", d.get("phone", HQ_PHONE))):
        if txt:
            s.pic(asset("icons", f"{icon}.png"), 90, y, 26, 26)
            s.text(132, y + 1, 1200, txt, size=22, color="E6E6E6", lh=1.2, wrap=False)
            y += 41
    return s


# --------------------------------------------------------------------------------------
# Public API
# --------------------------------------------------------------------------------------
def build_deck(spec: list[dict], out: str | os.PathLike, template: str | None = None,
               footer_text="Borek Solutions Group · boreksolutions.de · Confidential"):
    deck = Deck(template=template, footer_text=footer_text)
    for i, sl in enumerate(spec, 1):
        name = sl.get("layout")
        if name not in LAYOUTS:
            raise ValueError(f"slide {i}: unknown layout '{name}'. Available: {', '.join(sorted(LAYOUTS))}")
        LAYOUTS[name](deck, sl)
        deck.slides[-1].label = name
    return deck.save(out)


# Short field reference (also embedded into the LLM prompt of make_master_deck.py)
LAYOUT_DOCS = {
    "cover": 'F1 Cover. {kicker_right:"Family-owned since 1781", kicker:"Borek <unit> · <subtitle>", title:"<=2 lines, \\n between lines", '
             'intro:"<=3 lines", columns:[{title,text}] (0-3)}',
    "who_we_are": "F2 FIXED slide (history, map, management team, client logos). No content needed - add it unchanged after the cover.",
    "contrast": "L01 two cards + statement. {kicker,title,lead,left:{label,title,bullets:[4]},right:{label,title,bullets:[4]},"
                "statement:{kicker,text},note}",
    "four_cards": "L02 four problem cards + 3 facts. {kicker,title(1 line),lead,cards:[{title,text}x4],strip:{kicker,items:[{big,text}x3]}}",
    "matrix": 'L03 comparison matrix. {kicker,title,lead,columns:[<=5 provider names, LAST = Borek],rows:[{criterion,cells:["✓"|"—"|"short word"...]}] (<=8),statement}',
    "offers": "L04 three offers. {kicker,title(1 line),lead,cards:[{kicker,title,text,bullets:[3],price}x3],panel:{kicker,text},note}",
    "pillars": "L05 five pillars. {kicker,title,lead,cards:[{n:'01',title,text}x5],statement:{kicker,text}}",
    "divider": "L06 section divider. {kicker:'Section 01 · Topic',title(1 line),text(1 line)}",
    "process": "L07 three stages. {kicker,title(1 line),lead,stages:[{n,title,text,result_label,result}x3],panel:{kicker,text},note}",
    "matrix_notes": "L08 list + 2 notes. {kicker,title,lead,rows:[{code:'L1',title,text}x<=7],notes:[{title,text}x2]}",
    "data_table": "L10 table + points. {kicker,title,lead,table_label,headers:[4],rows:[[4 cells]...,last row = total],points:[{title,text}x3],callout}",
    "big_number": "L11 big number + bars. {kicker,title,lead,big:{label,value,text},bars_label,bars:[{title,text,pct 0-100,value}x3],note}",
    "flow_facts": "L12 flow + facts. {kicker,title,lead,steps:[5 short labels],facts:[{big,text}x4],note}",
    "faq": "L14 objections. {kicker,title(1 line),lead,items:[{q,a}x4],support}",
    "two_models": "L15 two models. {kicker,title,lead,models:[{kicker,title,text}x2],features:[{title,text}x4],note}",
    "phased": "L16 phased steps. {kicker,title(1 line),lead,steps:[{when:'Day 1-5',title,text}x4],panel:{kicker,text}}",
    "packages": "L17 packages. {kicker,title,lead,packages:[{title,sub,price,unit,bullets:[4]}x3],highlight:1,badge:'Popular',panel:{kicker,text}}",
    "capabilities": "L18 capabilities + stats. {kicker,title,lead,cards:[{title,text}x3],stats:[{big,text}x5],note}",
    "certificates": "L19 certificates + commitments. {kicker,title,lead,badges:[{title,text}x4],commitments:[{title,text}x6]}",
    "options": "L20 option comparison. {kicker,title,options:[name x3],rows:[{label,cells:[3],style:'bold'|'muted'|null}],availability_label,availability:[{text,on}x3],recommendation}",
    "price_tables": "L21 price tables. {kicker,title,lead,left:{kicker,headers:[3],rows:[[3]]},right:{kicker,headers:[5],rows:[[5]]},note}",
    "case_study": "L22 case study. {kicker,title,image:'path to screenshot 820x600',text,results:[{label,value,text}x2]}",
    "team": "L23 management team (max 8). {kicker,title,people:['Official Name',...]}",
    "locations": "L24 map. {kicker,title,stat:'250+',stat_label,locations:[[city,role]x4]}",
    "closing": "F3 closing. {title:'Let’s talk',contacts:['Official Name' x<=4],tagline,address,phone}",
}
