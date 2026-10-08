"""PowerPoint assembly: canonical deck (untouched) + appendix slides = Master Presentation.

The canonical package is opened and the appendix slides are added to it as new slide parts.
Nothing of the original slides is regenerated: their XML, relationships and media stay as
they are, and only the package-level lists (slide ids, relationships, content types) grow.
``verify_canonical_preserved`` proves that after every build. A slide looks the way it does
because of more than its own XML, so each original slide is compared together with everything
it depends on - layout, slide master, theme, notes, media, embedded objects - and the
presentation-level settings (slide size, default text styles, embedded fonts, table styles,
view and presentation properties) are compared as well, independent of zip-level details.
"""

from __future__ import annotations

import hashlib
import posixpath
import tempfile
import zipfile
from functools import lru_cache
from pathlib import Path
from typing import Any

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu

from services.presentation.borek_deck.engine import engine
from services.presentation.master_deck.layouts import NAVY, WHITE, Canvas, draw_slide
from services.presentation.master_deck.registry import MasterDeck, MasterDeckError, get_master, verify_master

_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
_R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
# Canonical slides the brand images of the appendix are taken from (same files, same look).
_LIGHT_SLIDE, _DIVIDER_SLIDE = 3, 8


class AssemblyError(RuntimeError):
    code = "MASTER_ASSEMBLY_FAILED"
    retryable = False


@lru_cache(maxsize=2)
def brand_images(master_id: str) -> dict[str, str]:
    """Logo and cover image exactly as the canonical deck embeds them, as files for the renderer."""
    master = get_master(master_id)
    deck = Presentation(str(master.pptx_path))
    pictures = lambda number: [shape for shape in deck.slides[number - 1].shapes if shape.shape_type == 13]  # noqa: E731
    light, divider = pictures(_LIGHT_SLIDE), pictures(_DIVIDER_SLIDE)
    if len(light) != 1 or len(divider) != 2:
        raise MasterDeckError(f"Master deck {master_id} does not carry the expected brand images")
    target = Path(tempfile.mkdtemp(prefix=f"{master_id}-brand-"))
    blobs = {"logo_dark": light[0].image.blob, "cover_bg": divider[0].image.blob, "logo_white": divider[1].image.blob}
    paths = {}
    for name, blob in blobs.items():
        path = target / f"{name}.png"
        path.write_bytes(blob)
        paths[name] = str(path)
    return paths


def _add_slide(deck: Any, layout: Any, *, dark: bool) -> Any:
    slide = deck.slides.add_slide(layout)
    for placeholder in list(slide.placeholders):
        placeholder._element.getparent().remove(placeholder._element)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(NAVY if dark else WHITE)
    return engine().bp.S(None, slide, dark, "none")


def _draw_all(deck: Any, layout: Any, contents: list[dict[str, Any]], master_id: str, *, first_page: int, total: int) -> list[str]:
    canvas = Canvas(brand_images(master_id))
    for offset, content in enumerate(contents):
        before = len(canvas.problems)
        slide = _add_slide(deck, layout, dark=content.get("layout") == "L06")
        draw_slide(slide, content, canvas, page=first_page + offset, total=total)
        for position in range(before, len(canvas.problems)):
            canvas.problems[position] = f"appendix slide {offset + 1} ({content.get('layout')}): {canvas.problems[position]}"
    return canvas.problems


def _blank_deck() -> tuple[Any, Any]:
    bp = engine().bp
    deck = Presentation()
    deck.slide_width, deck.slide_height = Emu(bp.SLIDE_W * bp.PX), Emu(bp.SLIDE_H * bp.PX)
    return deck, deck.slide_layouts[6]


def fit_problems(contents: list[dict[str, Any]], master_id: str | None = None) -> list[str]:
    """Lay the appendix out in memory; returns every text that does not fit its layout."""
    from services.presentation.master_deck.registry import DEFAULT_MASTER_ID

    deck, layout = _blank_deck()
    return _draw_all(deck, layout, contents, master_id or DEFAULT_MASTER_ID, first_page=1, total=len(contents))


def assemble(master: MasterDeck, contents: list[dict[str, Any]], out: Path) -> None:
    """Write the final deck: the verified canonical slides followed by the appendix slides."""
    verify_master(master)
    deck = Presentation(str(master.pptx_path))
    if len(deck.slides) != master.slide_count:
        raise MasterDeckError(f"Master deck {master.master_id} does not open with {master.slide_count} slides")
    total = master.slide_count + len(contents)
    problems = _draw_all(deck, deck.slide_layouts[0], contents, master.master_id, first_page=master.slide_count + 1, total=total)
    if problems:
        raise AssemblyError("; ".join(problems))
    deck.save(str(out))
    verify_canonical_preserved(master, out, expected_total=total)


# ---------------------------------------------------------------------------------- verification


def _slide_parts(package: zipfile.ZipFile) -> list[str]:
    """Slide part names in presentation order."""
    presentation = etree.fromstring(package.read("ppt/presentation.xml"))
    rels = etree.fromstring(package.read("ppt/_rels/presentation.xml.rels"))
    targets = {rel.get("Id"): rel.get("Target") for rel in rels.findall(f"{{{_REL_NS}}}Relationship")}
    ids = presentation.find(f"{{{_P_NS}}}sldIdLst")
    return [posixpath.normpath(posixpath.join("ppt", targets[item.get(f"{{{_R_NS}}}id")])) for item in ids]


def _canonical_xml(data: bytes) -> bytes:
    return etree.tostring(etree.fromstring(data), method="c14n")


def _rels_part(part: str) -> str:
    return posixpath.join(posixpath.dirname(part), "_rels", posixpath.basename(part) + ".rels")


class _Package:
    """Read access to a PowerPoint package for fingerprinting, with per-part caches."""

    def __init__(self, package: zipfile.ZipFile) -> None:
        self.zip = package
        self.names = set(package.namelist())
        types = etree.fromstring(package.read("[Content_Types].xml"))
        self._defaults = {item.get("Extension").lower(): item.get("ContentType") for item in types.findall(f"{{{_CT_NS}}}Default")}
        self._overrides = {item.get("PartName").lstrip("/"): item.get("ContentType") for item in types.findall(f"{{{_CT_NS}}}Override")}
        self._hashes: dict[str, bytes] = {}
        self._rels: dict[str, list[tuple[str, str, bool, str]]] = {}

    def content_type(self, part: str) -> str:
        return self._overrides.get(part) or self._defaults.get(part.rsplit(".", 1)[-1].lower(), "")

    def content_hash(self, part: str) -> bytes:
        """Checksum of a part's content type and content; XML is compared in canonical form."""
        if part not in self._hashes:
            data = self.zip.read(part)
            # PowerPoint parts are re-serialised when a deck is saved; everything else
            # (custom XML, media, embedded files) is copied and compared byte for byte.
            if part.startswith("ppt/") and part.endswith(".xml"):
                data = _canonical_xml(data)
            self._hashes[part] = hashlib.sha256(self.content_type(part).encode() + b"\0" + data).digest()
        return self._hashes[part]

    def relationships(self, part: str) -> list[tuple[str, str, bool, str]]:
        """(id, kind, external, target) of a part's relationships, ordered by id."""
        if part not in self._rels:
            found = []
            if _rels_part(part) in self.names:
                for rel in etree.fromstring(self.zip.read(_rels_part(part))).findall(f"{{{_REL_NS}}}Relationship"):
                    external = rel.get("TargetMode") == "External"
                    target = rel.get("Target")
                    if not external:
                        target = posixpath.normpath(posixpath.join(posixpath.dirname(part), target)).lstrip("/")
                    found.append((rel.get("Id"), rel.get("Type").rsplit("/", 1)[-1], external, target))
            self._rels[part] = sorted(found)
        return self._rels[part]

    def closure(self, part: str) -> list[tuple[str, str]]:
        """A part and everything it depends on, as (kind, checksum) in a deterministic order.

        Follows every internal relationship transitively (slide -> layout -> master -> theme,
        notes, media, charts, embedded objects ...). Parts are identified by the order they are
        reached in, never by name, so the result is the same when a package is re-saved with
        other part names - and different as soon as any dependency's content, content type or
        relationship changes.
        """
        order: dict[str, int] = {part: 0}
        kinds = {part: "slide"}
        queue = [part]
        entries: list[tuple[str, str]] = []
        while queue:
            current = queue.pop(0)
            if current not in self.names:
                raise AssemblyError(f"The deck references a part that is missing from the package ({kinds[current]})")
            digest = hashlib.sha256(self.content_hash(current))
            for rel_id, kind, external, target in self.relationships(current):
                if external:
                    digest.update(f"|{rel_id}|{kind}|external|{target}".encode())
                    continue
                if target not in order:
                    order[target] = len(order)
                    kinds[target] = kind
                    queue.append(target)
                digest.update(f"|{rel_id}|{kind}|{order[target]}".encode())
            entries.append((kinds[current], digest.hexdigest()))
        return entries


def slide_fingerprint(package: zipfile.ZipFile, part: str) -> str:
    """Checksum of a slide and of its complete dependency chain (see ``_Package.closure``)."""
    return _fingerprint(_Package(package).closure(part))


def _fingerprint(entries: list[tuple[str, str]]) -> str:
    return hashlib.sha256("\n".join(f"{kind}:{value}" for kind, value in entries).encode()).hexdigest()


def _slide_closures(path: Path) -> list[list[tuple[str, str]]]:
    with zipfile.ZipFile(path) as archive:
        package = _Package(archive)
        return [package.closure(part) for part in _slide_parts(archive)]


def canonical_fingerprints(path: Path) -> list[str]:
    return [_fingerprint(entries) for entries in _slide_closures(path)]


def presentation_settings(path: Path, *, slide_count: int) -> dict[str, str]:
    """Everything at presentation level that shapes how the first ``slide_count`` slides look.

    ``presentation.xml`` (slide size, default text styles, embedded fonts, master lists, and the
    ids of those slides - slides appended after them are ignored) and every non-slide part the
    presentation references: slide masters, notes master, theme, table styles, presentation and
    view properties, fonts, each with its own dependency chain.
    """
    with zipfile.ZipFile(path) as archive:
        package = _Package(archive)
        root = etree.fromstring(archive.read("ppt/presentation.xml"))
        ids = root.find(f"{{{_P_NS}}}sldIdLst")
        kept = {item.get(f"{{{_R_NS}}}id") for item in list(ids)[:slide_count]}
        for item in list(ids)[slide_count:]:
            ids.remove(item)
        settings = {"presentation.xml": hashlib.sha256(etree.tostring(root, method="c14n")).hexdigest()}
        for rel_id, kind, external, target in package.relationships("ppt/presentation.xml"):
            if kind == "slide":
                if rel_id in kept:  # slides appended after the canonical ones are not settings
                    settings[f"{kind} {rel_id}"] = "listed"
            elif external:
                settings[f"{kind} {rel_id}"] = f"external {target}"
            else:
                settings[f"{kind} {rel_id}"] = _fingerprint(package.closure(target))
        return settings


def _changed_dependency(before: list[tuple[str, str]], after: list[tuple[str, str]]) -> str:
    for (kind, old), (_kind, new) in zip(before, after, strict=False):
        if old != new:
            return kind
    return "dependencies"


def verify_canonical_preserved(master: MasterDeck, final_path: Path, *, expected_total: int) -> None:
    """The first slides of the final deck are the canonical slides, in order and unchanged.

    Unchanged includes everything they are rendered with: layouts, slide masters, themes, notes,
    media and the presentation-level settings.
    """
    expected = _slide_closures(master.pptx_path)
    actual = _slide_closures(final_path)
    if len(actual) != expected_total:
        raise AssemblyError(f"The assembled deck has {len(actual)} slides, expected {expected_total}")
    for number, (before, after) in enumerate(zip(expected, actual[: len(expected)], strict=True), start=1):
        if before != after:
            what = _changed_dependency(before, after)
            detail = "" if what == "slide" else f" (its {what})"
            raise AssemblyError(f"Canonical slide {number} was changed during assembly{detail}")
    before_settings = presentation_settings(master.pptx_path, slide_count=len(expected))
    after_settings = presentation_settings(final_path, slide_count=len(expected))
    if before_settings != after_settings:
        changed = sorted(
            key.split(" ")[0]
            for key in before_settings.keys() | after_settings.keys()
            if before_settings.get(key) != after_settings.get(key)
        )
        raise AssemblyError(f"The presentation settings of the canonical deck were changed during assembly ({', '.join(changed)})")
    try:
        reopened = Presentation(str(final_path))
    except Exception as exc:  # a package PowerPoint could not open is not a deck
        raise AssemblyError(f"The assembled deck cannot be opened: {exc}") from exc
    if len(reopened.slides) != expected_total:
        raise AssemblyError("The assembled deck does not reopen with the expected slide count")
