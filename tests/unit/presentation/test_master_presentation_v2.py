"""Master Presentation V2: frozen snapshot, manifest, post-meeting appendix planner, assembly."""

from __future__ import annotations

import copy
import zipfile
from pathlib import Path
from typing import Any

import pytest
from pptx import Presentation

from services.framework.discovery_analysis.pipeline import generate_discovery_analysis
from services.presentation.master_deck import appendix, appendix_v2, assembly, office, registry, v2
from services.presentation.master_deck.appendix import AppendixPlanError
from services.presentation.master_deck.plan import is_master_manifest, is_master_plan
from tests.fixtures.master_render_double import install_master_render_double

OPPORTUNITY = {
    "id": "11111111-1111-4111-8111-111111111112",
    "client_name": "Nordwind Maschinenbau",
    "opportunity_name": "AI introduction",
    "stage1_intake": {
        "poc_name": "Dana Weber",
        "sales_topic_description": "Faster quoting for the sales team",
        "about_company": "Family-owned machine builder with 420 employees.",
    },
}
EXCLUDED = "Send the pricing export by Friday."
OBSERVATION = "The service team could reuse the quoting assistant."
USE_CASE = {
    "fact_id": "reference.invoice-3way.delivery-pattern",
    "status": "resolved",
    "statement": "Invoice 3-way Match is typically delivered as a structured finance-operations engagement with a named delivery lead.",
    "service_key": "invoice_3way",
    "document_id": "REF-INVOICE-v1",
    "document_version": "1",
    "document_type": "reference",
    "corpus_id": "borek-internal",
    "corpus_version": "2026.09",
}


@pytest.fixture(autouse=True)
def _renderer(monkeypatch: pytest.MonkeyPatch) -> None:
    install_master_render_double(monkeypatch)


def finding(text: str, source: str = "transcript", status: str = "confirmed") -> dict[str, str]:
    return {"text": text, "source": source, "status": status}


def reviewed(**overrides: list[dict[str, str]]) -> dict[str, list[dict[str, str]]]:
    items = {
        "requirements": [finding("Quotes must go out within one day."), finding("Every quote needs a sign-off by the sales lead.")],
        "challenges": [finding("Pricing data sits in three systems.")],
        "priorities": [finding("Start with the sales team.", "both")],
        "opportunities": [finding(OBSERVATION, "personal_notes")],
        "discussed_solutions": [finding("A drafting assistant that prepares the quote from the request.")],
        "decisions": [finding("Run a pilot with ten quotes.")],
        "follow_ups": [finding(EXCLUDED, status="excluded")],
    }
    items.update(overrides)
    return items


def approved_row(**overrides: Any) -> dict[str, Any]:
    paper = generate_discovery_analysis(
        OPPORTUNITY, persist=lambda _paper: None, generated_at="2026-10-08T10:00:00Z", document_id="22222222-2222-4222-8222-222222222222"
    )
    row = {
        "id": "33333333-3333-4333-8333-333333333333",
        "opportunity_id": OPPORTUNITY["id"],
        "document_id": paper["document_id"],
        "version_number": 1,
        "status": "approved",
        "paper_json": paper,
    }
    row.update(overrides)
    return row


def snapshot_for(items: dict[str, list[dict[str, str]]] | None = None, **overrides: Any) -> dict[str, Any]:
    arguments: dict[str, Any] = {
        "opportunity_id": OPPORTUNITY["id"],
        "presentation_id": "44444444-4444-4444-8444-444444444444",
        "base_presentation_version_id": "55555555-5555-4555-8555-555555555555",
        "approved_discovery": approved_row(),
        "transcript": {"transcript_id": "66666666-6666-4666-8666-666666666666", "revision": "a" * 64, "file_name": "meeting.txt"},
        "extraction": {"generated_at": "2026-10-08T11:00:00Z", "execution_mode": "fixture", "personal_notes_updated_at": "2026-10-08T10:30:00Z"},
        "personal_notes": {"updated_at": "2026-10-08T10:30:00Z", "text": f"Opportunity: {OBSERVATION}"},
        "reviewed_items": items or reviewed(),
        "use_cases": [copy.deepcopy(USE_CASE)],
        "meeting_review": {
            "confirmed_at": "2026-10-08T11:05:00Z",
            "confirmed_by": "aaaaaaaa-aaaa-4aaa-8aaa-aaaaaaaaaaaa",
            "review_fingerprint": "b" * 64,
            "source_hash": "c" * 64,
        },
    }
    arguments.update(overrides)
    return v2.build_snapshot(**arguments)


def appendix_of(snapshot: dict[str, Any]) -> list[dict[str, Any]]:
    return appendix_v2.plan_v2_appendix(snapshot, first_order=27)


def texts(slides: list[dict[str, Any]]) -> list[str]:
    return [text for slide in slides for text in appendix._texts(slide["content"])]


def test_snapshot_freezes_content_and_identities_and_keeps_excluded_findings_out() -> None:
    snapshot = snapshot_for()
    assert (snapshot["kind"], snapshot["planner_version"]) == ("master_presentation_v2_snapshot", "master-presentation-v2:v2")
    assert snapshot["master"]["master_sha256"] == "2c23670e46dc2ce71d3527b0acc846cad54572d9817977ec26c3efe17f51dd5a"
    assert snapshot["approved_discovery"]["version_id"] == "33333333-3333-4333-8333-333333333333"
    assert snapshot["approved_discovery"]["analysis"]["areas"], "the Discovery content itself is frozen, not only its id"
    assert snapshot["personal_notes"] == {"updated_at": "2026-10-08T10:30:00Z", "text": f"Opportunity: {OBSERVATION}"}
    assert snapshot["findings"]["follow_ups"] == [] and snapshot["excluded_findings"] == [EXCLUDED]
    assert {"text": OBSERVATION, "source": "personal_notes"} in snapshot["findings"]["opportunities"]
    assert snapshot["use_cases"][0]["corpus_version"] == "2026.09"
    assert v2.snapshot_hash(snapshot) == v2.snapshot_hash(snapshot_for()), "the same sources freeze to the same snapshot"

    manifest = v2.generation_manifest(snapshot)
    assert (manifest["kind"], manifest["product_version"], manifest["product_stage"]) == ("master_presentation_v2", "V2", "post_meeting")
    assert manifest["snapshot_hash"] == v2.snapshot_hash(snapshot) and manifest["source_hash"] == "c" * 64
    assert (manifest["confirmed_finding_count"], manifest["excluded_finding_count"]) == (7, 1)
    assert is_master_manifest(manifest) and "Quotes must" not in str(manifest)
    assert v2.same_source(manifest, v2.generation_manifest(snapshot_for()))
    assert not v2.same_source(manifest, {**manifest, "kind": "master_presentation_v1"})
    assert not v2.same_source(manifest, {key: value for key, value in manifest.items() if key != "generation_fingerprint"})


def test_the_generation_fingerprint_follows_the_inputs_not_the_time_of_confirmation() -> None:
    snapshot = snapshot_for()
    manifest = v2.generation_manifest(snapshot)
    fingerprint = manifest["generation_fingerprint"]
    assert len(fingerprint) == 64 and fingerprint == v2.generation_fingerprint(snapshot_for())

    # Confirmed again, analysed again, notes saved again with the same text: the same generation input.
    renewed = snapshot_for(
        meeting_review={
            "confirmed_at": "2026-10-09T08:00:00Z",
            "confirmed_by": "bbbbbbbb-bbbb-4bbb-8bbb-bbbbbbbbbbbb",
            "review_fingerprint": "d" * 64,
            "source_hash": "e" * 64,
        },
        extraction={"generated_at": "2026-10-09T07:55:00Z", "execution_mode": "fixture", "personal_notes_updated_at": "2026-10-09T07:50:00Z"},
        personal_notes={"updated_at": "2026-10-09T07:50:00Z", "text": f"Opportunity: {OBSERVATION}"},
    )
    renewed_manifest = v2.generation_manifest(renewed)
    assert renewed_manifest["generation_fingerprint"] == fingerprint and v2.same_source(manifest, renewed_manifest)
    # Each manifest still records its own confirmation: traceability is kept, lineage is not rewritten.
    assert (manifest["meeting_review_confirmed_at"], renewed_manifest["meeting_review_confirmed_at"]) == ("2026-10-08T11:05:00Z", "2026-10-09T08:00:00Z")
    assert manifest["snapshot_hash"] != renewed_manifest["snapshot_hash"] and manifest["source_hash"] != renewed_manifest["source_hash"]

    row = approved_row()
    other_discovery = {**row, "id": "77777777-7777-4777-8777-777777777777"}
    changed_content = {**row, "paper_json": {**row["paper_json"], "presentation_brief": {**row["paper_json"]["presentation_brief"], "core_thesis": "Working hypothesis: another thesis."}}}
    changes = {
        "a finding was added": {"reviewed_items": reviewed(decisions=[finding("Run a pilot with ten quotes."), finding("Review after six weeks.")])},
        "a finding was reworded": {"reviewed_items": reviewed(challenges=[finding("Pricing data sits in four systems.")])},
        "an exclusion was lifted": {"reviewed_items": reviewed(follow_ups=[finding(EXCLUDED)])},
        "another finding was excluded": {"reviewed_items": reviewed(decisions=[finding("Run a pilot with ten quotes.", status="excluded")])},
        "a finding changed its source": {"reviewed_items": reviewed(priorities=[finding("Start with the sales team.", "transcript")])},
        "another Discovery version was approved": {"approved_discovery": other_discovery},
        "the approved Discovery content differs": {"approved_discovery": changed_content},
        "the notes text changed": {"personal_notes": {"updated_at": "2026-10-08T10:30:00Z", "text": "Something else entirely."}},
        "the transcript content changed": {"transcript": {"transcript_id": "66666666-6666-4666-8666-666666666666", "revision": "f" * 64, "file_name": "meeting.txt"}},
        "another transcript was analysed": {"transcript": {"transcript_id": "88888888-8888-4888-8888-888888888888", "revision": "a" * 64, "file_name": "meeting.txt"}},
        "the use cases changed": {"use_cases": []},
        "the reference was revised": {"use_cases": [{**USE_CASE, "corpus_version": "2026.10"}]},
        "V1 was regenerated": {"base_presentation_version_id": "99999999-9999-4999-8999-999999999999"},
    }
    seen = {fingerprint}
    for reason, overrides in changes.items():
        other = v2.generation_manifest(snapshot_for(**overrides))
        assert not v2.same_source(manifest, other), reason
        seen.add(other["generation_fingerprint"])
    assert len(seen) == len(changes) + 1, "every kind of change gives its own fingerprint"


@pytest.mark.parametrize(
    ("overrides", "message"),
    [
        ({"approved_discovery": approved_row(status="draft")}, "approved Discovery version"),
        ({"approved_discovery": approved_row(opportunity_id="99999999-9999-4999-8999-999999999999")}, "another opportunity"),
        ({"approved_discovery": {**approved_row(), "paper_json": {"schema_version": "1.0"}}}, "schema 2.0"),
        ({"use_cases": [{**USE_CASE, "status": "unresolved"}]}, "no longer available"),
        ({"reviewed_items": reviewed(decisions=[finding("Budget approved.", "unverified")])}, "no verified source"),
    ],
)
def test_incomplete_or_foreign_sources_are_refused(overrides: dict[str, Any], message: str) -> None:
    with pytest.raises(v2.SnapshotError, match=message):
        snapshot_for(**overrides)


def test_the_appendix_is_built_from_the_meeting_and_keeps_every_origin_visible() -> None:
    snapshot = snapshot_for()
    slides = appendix_of(snapshot)
    assert [slide["order"] for slide in slides] == list(range(27, 27 + len(slides)))
    assert {slide["layout_id"] for slide in slides} <= set(appendix_v2.PLANNER_LAYOUTS)
    by_kicker = {slide["content"]["kicker"]: slide for slide in slides}
    assert slides[0]["content"]["kicker"] == "Appendix · After our first meeting"
    assert slides[0]["content"]["title"] == "Nordwind Maschinenbau"

    heard = by_kicker["From the meeting · Client statements"]["content"]
    assert [(row["code"], row["title"], row["text"]) for row in heard["rows"]] == [
        ("R1", "Requirement", "Quotes must go out within one day."),
        ("R2", "Requirement", "Every quote needs a sign-off by the sales lead."),
        ("P1", "Priority", "Start with the sales team."),
        ("C1", "Challenge", "Pricing data sits in three systems."),
        ("S1", "Discussed solution", "A drafting assistant that prepares the quote from the request."),
    ]
    # One observation and one reference share a page: both are Borek's content, labelled row by row.
    ours = by_kicker["From Borek · Not client statements"]["content"]
    assert [(row["code"], row["title"], row["text"]) for row in ours["rows"]] == [
        ("N1", "Our observation", OBSERVATION),
        ("B1", "Borek reference", USE_CASE["statement"]),
    ]
    assert "None of this was said by Nordwind Maschinenbau" in ours["lead"]
    assert [note["title"] for note in ours["notes"]] == ["Our observations", "Borek references"]

    closing = by_kicker["Decisions · Next steps · Open questions"]["content"]
    assert ("D1", "Decision", "Run a pilot with ten quotes.") in [(row["code"], row["title"], row["text"]) for row in closing["rows"]]
    questions = [row["text"] for row in closing["rows"] if row["title"] == "Open question"]
    analysis = snapshot["approved_discovery"]["analysis"]
    asked = {q for area in analysis["areas"] for item in area["opportunities"] for q in item["discovery_questions"]}
    assert 1 <= len(questions) <= 3 and set(questions) <= asked, "open questions are the Discovery analysis' own questions"
    assert all(question.endswith("?") and len(question.split()) >= 5 for question in questions), "never a bare topic label"
    assert not set(questions) & set(analysis["research"]["unknown_facts"])

    approach = slides[-1]["content"]
    assert approach["kicker"].startswith("Borek recommendation") and approach["statement"] == {
        "kicker": "Decided in the meeting", "text": "Run a pilot with ten quotes.",
    }

    everything = texts(slides)
    assert EXCLUDED not in everything
    # Every confirmed finding is on exactly one slide, word for word; every slide cites its sources.
    for item in appendix_v2.confirmed_findings(snapshot):
        holders = [slide for slide in slides if item["text"] in [row["text"] for row in slide["content"].get("rows", [])]]
        assert len(holders) == 1, item
        borek_page = "Not client statements" in holders[0]["content"]["kicker"]
        assert borek_page == (item["source"] == "personal_notes"), item
    for slide in slides:
        assert slide["source_references"] and all(appendix_v2.resolve_v2_reference(snapshot, ref) is not None for ref in slide["source_references"])
    # It is not the V1 appendix under a new name.
    v1 = appendix.plan_appendix(approved_row()["paper_json"], first_order=27)
    assert [slide["content"]["title"] for slide in slides] != [slide["content"]["title"] for slide in v1]
    assert not {slide["content"]["kicker"] for slide in slides} & {slide["content"]["kicker"] for slide in v1}


def test_the_slide_plan_follows_the_amount_of_confirmed_information() -> None:
    minimal = appendix_of(snapshot_for(reviewed(
        requirements=[finding("Quotes must go out within one day.")], challenges=[], priorities=[], opportunities=[],
        discussed_solutions=[], decisions=[], follow_ups=[],
    ), use_cases=[]))
    kickers = [slide["content"]["kicker"] for slide in minimal]
    assert not any("Not client statements" in kicker for kicker in kickers), "nothing from Borek's side to show"
    assert kickers.count("From the meeting · Client statements") == 1
    assert minimal[-1]["content"]["statement"]["kicker"] == "Borek recommendation", "no decision: nothing is presented as decided"

    many = reviewed(
        requirements=[finding(f"Requirement number {index} for the quoting process must be met by the sales office.") for index in range(1, 9)],
        challenges=[finding(f"Challenge number {index} slows the team down every week.") for index in range(1, 6)],
        opportunities=[finding(f"Observation number {index} about the way the office handles requests.", "personal_notes") for index in range(1, 7)],
        decisions=[finding("Run a pilot with ten quotes."), finding("Review the pilot after six weeks.")],
        follow_ups=[finding("Send the pricing export by Friday.")],
    )
    rich = appendix_of(snapshot_for(many))
    titles = [slide["content"]["title"] for slide in rich]
    assert "What you need and what comes first (1/2)" in titles and "What you need and what comes first (2/2)" in titles
    assert "Where it is difficult today and where you see potential" in titles and "Solutions we talked about" in titles
    # Six observations and a reference do not fit one page: each kind gets its own, still labelled.
    assert "What we noted for ourselves" in titles and "Borek references that fit what we heard" in titles
    assert "What we add from our side" not in titles
    assert len(rich) > len(minimal) + 4
    assert all(len(slide["content"].get("rows", [])) <= 6 for slide in rich)
    assert not appendix.layout_fit_problems(rich)
    # A page with a single row only exists where the meeting gave a single statement of that kind.
    single = [slide for slide in appendix_of(snapshot_for()) if len(slide["content"].get("rows", [])) == 1]
    assert single == [], "the short reference meeting has no single-row page"
    # A lone observation (no reference selected) opens the recommendation page instead of filling one alone.
    lone = appendix_of(snapshot_for(use_cases=[]))
    (ours,) = [slide for slide in lone if "Not client statements" in slide["content"]["kicker"]]
    assert ours["content"]["title"] == "What we add and where we would start"
    codes = [row["code"] for row in ours["content"]["rows"]]
    assert codes[0] == "N1" and len(codes) > 1 and all(code.startswith("A") for code in codes[1:])
    assert ours["content"]["rows"][0]["title"] == "Our observation"
    assert [note["title"] for note in ours["content"]["notes"]] == ["Our observations", "Why these starting points"]
    assert any(reference.startswith("finding:") for reference in ours["source_references"]), "the recommendations still cite their evidence"
    assert "benchmarks, not commitments" in ours["content"]["lead"]
    assert all(slide["content"]["kicker"] != "Borek recommendation" for slide in lone), "no second, separate recommendation page"
    assert not [slide for slide in lone if len(slide["content"].get("rows", [])) == 1]


def test_recommendations_need_clear_support_in_what_the_client_said() -> None:
    snapshot = snapshot_for()
    slides = appendix_of(snapshot)
    (recommendation,) = [slide for slide in slides if slide["content"]["kicker"] == "Borek recommendation"]
    rows = recommendation["content"]["rows"]
    analysis = snapshot["approved_discovery"]["analysis"]
    results = {item["title"]: item["result"] for area in analysis["areas"] for item in area["opportunities"]}
    assert 1 <= len(rows) <= 4 and all(results[row["title"]] == row["text"] for row in rows), "taken from the approved analysis as they are"
    assert rows[0]["title"] == "Inquiry intake & quote drafting", "the best supported opportunity comes first"
    # The page names the client statements it rests on, in its references and in plain words.
    evidence = [reference for reference in recommendation["source_references"] if reference.startswith("finding:")]
    assert evidence and all(appendix_v2.resolve_v2_reference(snapshot, reference)["source"] in appendix_v2.CLIENT_SOURCES for reference in evidence)
    why = recommendation["content"]["notes"][0]
    assert why["title"] == "Why these" and "you described in our first meeting" in why["text"] and "for you to confirm" in why["text"]
    assert recommendation["content"]["notes"][1]["text"] == appendix.BENCHMARK_NOTE

    # One shared word is a coincidence: every recommended opportunity shares at least two subject words.
    matches = appendix_v2._matching_opportunities(analysis, appendix_v2.confirmed_findings(snapshot))
    recommended = appendix_v2._recommended(matches)
    assert [match["item"]["title"] for match in recommended] == [row["title"] for row in rows]
    assert all(len(match["shared_terms"]) >= 2 for match in recommended)
    assert not {"every", "system", "team", "data"} & {term for match in matches for term in match["shared_terms"]}, "generic words carry no support"
    assert appendix_v2._keywords("Quotes and quoting processes for every team") == {"quote", "quot"} or "quote" in appendix_v2._keywords("Quotes for every team")

    # Nothing related was said: no recommendation page and no open questions, rather than weak ones.
    unrelated = reviewed(
        requirements=[finding("Zebra crossing paint.")], challenges=[], priorities=[], discussed_solutions=[], decisions=[],
        opportunities=[finding("Quotes, pricing and the sales pipeline could all be automated.", "personal_notes")],
    )
    plan = appendix_of(snapshot_for(unrelated))
    assert all(slide["content"]["kicker"] != "Borek recommendation" for slide in plan), "an owner note alone recommends nothing"
    assert all(row["title"] != "Open question" for slide in plan for row in slide["content"].get("rows", []))
    assert any("Not client statements" in slide["content"]["kicker"] for slide in plan)
    # One weakly related statement gives a question to ask, not a recommendation.
    weak = appendix_of(snapshot_for(reviewed(
        requirements=[finding("We lose many bids.")], challenges=[], priorities=[], opportunities=[], discussed_solutions=[], decisions=[],
    ), use_cases=[]))
    assert all(slide["content"]["kicker"] != "Borek recommendation" for slide in weak)


def test_nothing_confirmed_nothing_invented_and_nothing_cut() -> None:
    nothing = reviewed(**{category: [] for category in appendix_v2.CATEGORIES})
    with pytest.raises(AppendixPlanError, match="at least one confirmed meeting finding"):
        appendix_of(snapshot_for(nothing))

    snapshot = snapshot_for()
    slides = appendix_of(snapshot)
    invented = copy.deepcopy(slides)
    invented[1]["content"]["rows"][0]["text"] = "Quotes must go out within 37 hours and save 250000 EUR."
    with pytest.raises(AppendixPlanError, match="figures that are not in the frozen sources: 250000, 37"):
        appendix_v2.validate_v2_appendix(invented, snapshot, first_order=27)
    leaked = copy.deepcopy(slides)
    leaked[1]["content"]["rows"][0]["text"] = EXCLUDED
    with pytest.raises(AppendixPlanError, match="uses a finding the owner excluded"):
        appendix_v2.validate_v2_appendix(leaked, snapshot, first_order=27)
    relabelled = copy.deepcopy(slides)
    relabelled[1]["content"]["rows"][0]["text"] = OBSERVATION
    with pytest.raises(AppendixPlanError, match="presents an owner observation as a client statement"):
        appendix_v2.validate_v2_appendix(relabelled, snapshot, first_order=27)
    borek_index = next(index for index, slide in enumerate(slides) if "Not client statements" in slide["content"]["kicker"])
    misplaced = copy.deepcopy(slides)
    misplaced[borek_index]["content"]["rows"][0]["text"] = "Pricing data sits in three systems."
    with pytest.raises(AppendixPlanError, match="presents a client statement as Borek's own content"):
        appendix_v2.validate_v2_appendix(misplaced, snapshot, first_order=27)
    vague = copy.deepcopy(slides)
    question_slide = next(slide for slide in vague if any(row["title"] == "Open question" for row in slide["content"].get("rows", [])))
    next(row for row in question_slide["content"]["rows"] if row["title"] == "Open question")["text"] = "Website"
    with pytest.raises(AppendixPlanError, match="lists a topic label as an open question: 'Website'"):
        appendix_v2.validate_v2_appendix(vague, snapshot, first_order=27)
    unsupported = copy.deepcopy(slides)
    recommendation = next(slide for slide in unsupported if slide["content"]["kicker"] == "Borek recommendation")
    recommendation["source_references"] = [reference for reference in recommendation["source_references"] if not reference.startswith("finding:")]
    with pytest.raises(AppendixPlanError, match="recommends without naming the client statements it rests on"):
        appendix_v2.validate_v2_appendix(unsupported, snapshot, first_order=27)
    dropped = [slide for slide in copy.deepcopy(slides) if "Not client statements" not in slide["content"]["kicker"]]
    for index, slide in enumerate(dropped):
        slide["order"] = 27 + index
    with pytest.raises(AppendixPlanError, match="neither shown nor reported as omitted"):
        appendix_v2.validate_v2_appendix(dropped, snapshot, first_order=27)

    # A statement too long for a slide is left out and reported, never shortened.
    essay = "The client explained at length that " + "the current quoting process depends on several people and tools " * 30
    long_plan = v2.build_plan(snapshot_for(reviewed(challenges=[finding(essay.strip() + ".")])))
    assert long_plan["appendix"]["omitted_findings"] == ["finding:challenges:0"]
    assert all("depends on several people" not in text for slide in long_plan["slides"][26:] for text in appendix._texts(slide["borekSlide"]))


def test_plan_and_assembly_keep_the_canonical_deck_and_only_accept_their_own_snapshot(tmp_path: Path) -> None:
    snapshot = snapshot_for()
    plan = v2.build_plan(snapshot)
    manifest = v2.generation_manifest(snapshot)
    assert (plan["deck_kind"], plan["product_version"]) == ("master_v2", "V2") and is_master_plan(plan)
    assert [slide["layoutId"] for slide in plan["slides"][:26]] == ["CANONICAL"] * 26
    assert plan["appendix"] == {
        "first_slide": 27, "slide_count": len(plan["slides"]) - 26, "source_hash": manifest["snapshot_hash"], "omitted_findings": [],
    }
    v2.verify_frozen_job(manifest=manifest, snapshot=snapshot, plan_json=plan)
    altered = copy.deepcopy(snapshot)
    altered["findings"]["decisions"].append({"text": "The budget is approved.", "source": "transcript"})
    for arguments, message in (
        ({"manifest": manifest, "snapshot": altered, "plan_json": plan}, "does not match the generation manifest"),
        ({"manifest": manifest, "snapshot": None, "plan_json": plan}, "was not stored on this job"),
        ({"manifest": {**manifest, "source_hash": "e" * 64}, "snapshot": snapshot, "plan_json": plan}, "does not describe the frozen source snapshot"),
        ({"manifest": manifest, "snapshot": snapshot, "plan_json": {**plan, "deck_kind": "master_v1"}}, "no longer points at a Master Presentation V2 plan"),
        ({"manifest": manifest, "snapshot": snapshot, "plan_json": v2.build_plan(altered)}, "built from other sources"),
        ({"manifest": {**manifest, "kind": "master_presentation_v1"}, "snapshot": snapshot, "plan_json": plan}, "no Master Presentation V2 manifest"),
    ):
        with pytest.raises(v2.SnapshotError, match=message):
            v2.verify_frozen_job(**arguments)
    with pytest.raises(v2.SnapshotError, match="different canonical master deck"):
        v2.build_plan({**snapshot, "master": {**snapshot["master"], "master_sha256": "0" * 64}})

    master = registry.get_master()
    final = tmp_path / "v2.pptx"
    contents = [slide["borekSlide"] for slide in plan["slides"][26:]]
    assembly.assemble(master, contents, final)
    prints = assembly.canonical_fingerprints(final)
    assert len(prints) == 26 + len(contents) and prints[:26] == assembly.canonical_fingerprints(master.pptx_path)
    assert assembly.presentation_settings(final, slide_count=26) == assembly.presentation_settings(master.pptx_path, slide_count=26)
    assert registry.file_sha256(master.pptx_path) == master.sha256
    deck = Presentation(str(final))
    total = len(deck.slides)
    last = [shape.text_frame.text for shape in deck.slides[total - 1].shapes if shape.has_text_frame]
    assert f"{total:02d} / {total:02d}" in last
    with zipfile.ZipFile(final) as package:
        for part in assembly._slide_parts(package)[26:]:
            xml = package.read(part).decode("utf-8")
            assert "roundRect" not in xml and "gradFill" not in xml and 'typeface="Inter"' in xml


real_renderer = pytest.mark.skipif(not office.renderer_available(), reason="LibreOffice/poppler are not installed here")


@real_renderer
@pytest.mark.parametrize("case", ["short", "long"])
def test_libreoffice_renders_the_v2_deck_in_inter_without_overlapping_text(case: str, tmp_path: Path) -> None:
    import subprocess
    from xml.etree import ElementTree

    items = reviewed() if case == "short" else reviewed(
        requirements=[finding(f"Requirement number {index} for the quoting process must be met by the sales office.") for index in range(1, 9)],
        challenges=[finding(f"Challenge number {index} slows the team down every week and costs the sales office time it does not have.") for index in range(1, 6)],
    )
    contents = [slide["content"] for slide in appendix_of(snapshot_for(items))]
    deck = tmp_path / "deck.pptx"
    assembly.assemble(registry.get_master(), contents, deck)
    pdf, previews = office.render_pptx(deck, tmp_path, expected_slides=26 + len(contents))
    assert len(previews) == 26 + len(contents)
    assert all(name.lower().startswith("inter") for name in office.pdf_fonts(pdf))
    layout = tmp_path / "layout.html"
    subprocess.run(["pdftotext", "-bbox-layout", "-f", "27", str(pdf), str(layout)], check=True, timeout=120)
    pages = [node for node in ElementTree.parse(layout).getroot().iter() if node.tag.endswith("page")]
    assert len(pages) == len(contents)
    for number, page in enumerate(pages, start=27):
        lines = [
            (float(n.get("xMin")), float(n.get("yMin")), float(n.get("xMax")), float(n.get("yMax")), " ".join(w.text or "" for w in n))
            for n in page.iter()
            if n.tag.endswith("line")
        ]
        assert lines, f"slide {number} has no text"
        assert all(y1 <= float(page.get("height")) and x1 <= float(page.get("width")) for _x0, _y0, x1, y1, _t in lines), f"text runs off slide {number}"
        for index, (ax0, ay0, ax1, ay1, text_a) in enumerate(lines):
            for bx0, by0, bx1, by1, text_b in lines[index + 1 :]:
                overlap_x = min(ax1, bx1) - max(ax0, bx0)
                overlap_y = min(ay1, by1) - max(ay0, by0)
                assert not (overlap_x > 2 and overlap_y > 0.35 * min(ay1 - ay0, by1 - by0)), f"slide {number}: '{text_a[:40]}' overlaps '{text_b[:40]}'"
