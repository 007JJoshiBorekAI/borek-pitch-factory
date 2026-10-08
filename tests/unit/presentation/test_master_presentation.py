"""Master Presentation V1: canonical master integrity, appendix planning, assembly and rendering."""

from __future__ import annotations

import copy
import hashlib
import io
import json
import re
import shutil
import zipfile
from pathlib import Path
from typing import Any

import pytest
from pptx import Presentation

from services.framework.discovery_analysis.pipeline import finalize, generate_discovery_analysis
from services.presentation.borek_deck.deck_plan import borek_slide_spec
from services.presentation.borek_deck.render import BorekDeckRenderError, render_deck_bundle
from services.presentation.master_deck import appendix, assembly, office, registry
from services.presentation.master_deck.layouts import LAYOUT_IDS, LayoutError, implemented_layouts
from services.presentation.master_deck.plan import (
    CANONICAL_LAYOUT,
    MASTER_V1,
    build_master_plan,
    generation_manifest,
    same_source,
)
from services.presentation.master_deck.render import render_master_bundle
from tests.fixtures.master_render_double import install_master_render_double

REAL_RENDER_PPTX = office.render_pptx  # the production function, before any test double
SUPPLIED_SHA256 = "2c23670e46dc2ce71d3527b0acc846cad54572d9817977ec26c3efe17f51dd5a"
MINIMAL = {"id": "11111111-1111-4111-8111-111111111111", "client_name": "", "opportunity_name": "", "stage1_intake": None}
MEDIUM = {
    "id": "11111111-1111-4111-8111-111111111112",
    "client_name": "Nordwind Maschinenbau",
    "opportunity_name": "AI introduction",
    "stage1_intake": {
        "poc_name": "Dana Weber",
        "sales_topic_description": "Faster quoting for the sales team",
        "about_company": "Family-owned machine builder with 420 employees.",
    },
}
RICH = {
    "id": "11111111-1111-4111-8111-111111111113",
    "client_name": "Hansa Logistik Gruppe",
    "opportunity_name": "Automation programme",
    "stage1_intake": {
        "client_web_page": "https://hansa.example",
        "poc_name": "Jonas Brandt",
        "sales_topic_description": "Automate order intake, invoice processing and supplier documents across three sites",
        "about_company": "Logistics group with 1,200 employees. Quality audits and onboarding of new planners bind senior staff; capacity planning runs in spreadsheets.",
    },
}
CASES = {"minimal": MINIMAL, "medium": MEDIUM, "rich": RICH}


@pytest.fixture(autouse=True)
def _renderer(monkeypatch: pytest.MonkeyPatch) -> None:
    install_master_render_double(monkeypatch)


def paper_for(opportunity: dict[str, Any]) -> dict[str, Any]:
    return generate_discovery_analysis(
        opportunity,
        persist=lambda _paper: None,
        generated_at="2026-10-08T10:00:00Z",
        document_id="22222222-2222-4222-8222-222222222222",
    )


def approved_row(opportunity: dict[str, Any] = MEDIUM, **overrides: Any) -> dict[str, Any]:
    paper = paper_for(opportunity)
    row = {
        "id": "33333333-3333-4333-8333-333333333333",
        "opportunity_id": opportunity["id"],
        "document_id": paper["document_id"],
        "status": "approved",
        "paper_json": paper,
    }
    row.update(overrides)
    return row


def specs_of(plan: dict[str, Any]) -> list[dict[str, Any]]:
    return [borek_slide_spec(planned) for planned in plan["slides"]]


def layouts_of(opportunity: dict[str, Any]) -> list[str]:
    return [slide["layout_id"] for slide in appendix.plan_appendix(paper_for(opportunity), first_order=27)]


# --- master integrity ---------------------------------------------------------------------


def test_canonical_master_is_the_supplied_deck() -> None:
    master = registry.get_master()
    assert (master.master_id, master.master_version, master.slide_count, master.language) == (
        "borek_ai_tech_en_v1", "1.0", 26, "en"
    )
    assert hashlib.sha256(master.pptx_path.read_bytes()).hexdigest() == master.sha256 == SUPPLIED_SHA256
    manifest = registry.verify_master(master)
    assert [slide["number"] for slide in manifest["slides"]] == list(range(1, 27))
    assert len(Presentation(str(master.pptx_path)).slides) == 26
    titles = registry.canonical_slide_titles()
    assert titles[0] == "Your AI department. Delivered, not built."
    assert titles[7] == "Borek AI Suite" and titles[16] == "Borek AI Department" and titles[18] == "Borek AI Projects"
    assert titles[24].startswith("One rate card, one tier table") and titles[25] == "Let’s talk"
    # Only the deck and its manifest are stored: previews and PDF are rendered from the assembled deck.
    assert sorted(path.name for path in master.directory.iterdir()) == ["deck.pptx", "manifest.json"]
    for unknown in ("../../etc/passwd", "some_other_deck", "C:/decks/other.pptx"):
        with pytest.raises(registry.MasterDeckError, match="Unknown master deck"):
            registry.get_master(unknown)


def test_replaced_corrupted_or_missing_master_is_rejected(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    master = registry.get_master()
    specs = specs_of(build_master_plan(approved_row()))  # planned while the master is intact
    shutil.copytree(master.directory, tmp_path / master.master_id)
    monkeypatch.setattr(registry, "ASSETS", tmp_path)
    assert registry.verify_master(master)["slide_count"] == 26

    stored = tmp_path / master.master_id / "deck.pptx"
    original = stored.read_bytes()
    stored.write_bytes(original + b"tampered")
    with pytest.raises(registry.MasterDeckError, match="does not match its registered checksum"):
        registry.verify_master(master)
    with pytest.raises(BorekDeckRenderError, match="registered checksum"):
        render_deck_bundle(specs, deck_kind=MASTER_V1)

    # Another valid deck (e.g. a fixture) is not accepted in its place.
    other = Presentation()
    other.slides.add_slide(other.slide_layouts[6])
    other.save(str(stored))
    with pytest.raises(registry.MasterDeckError, match="registered checksum"):
        registry.verify_master(master)

    stored.unlink()
    with pytest.raises(registry.MasterDeckError, match="is missing"):
        registry.verify_master(master)

    stored.write_bytes(original)
    (tmp_path / master.master_id / "manifest.json").unlink()
    with pytest.raises(registry.MasterDeckError, match="has no manifest"):
        registry.verify_master(master)


# --- Discovery source ---------------------------------------------------------------------


def test_only_an_approved_discovery_v2_version_is_accepted() -> None:
    row = approved_row()
    assert generation_manifest(row) == {
        "schema_version": "1.0",
        "kind": "master_presentation_v1",
        "product_version": "V1",
        "product_stage": "pre_meeting",
        "master_id": "borek_ai_tech_en_v1",
        "master_version": "1.0",
        "master_sha256": SUPPLIED_SHA256,
        "master_slide_count": 26,
        "approved_discovery_version_id": row["id"],
        "approved_discovery_document_id": row["document_id"],
        "discovery_schema_version": "2.0",
        "appendix_source_hash": appendix.appendix_source_hash(row["paper_json"]),
    }
    with pytest.raises(appendix.AppendixPlanError, match="approved Discovery version"):
        generation_manifest(approved_row(status="draft"))
    v1 = approved_row()
    v1["paper_json"] = {"schema_version": "1.0", "pages": []}
    with pytest.raises(appendix.AppendixPlanError, match="schema 2.0"):
        build_master_plan(v1)
    no_brief = approved_row()
    no_brief["paper_json"]["presentation_brief"] = None
    with pytest.raises(appendix.AppendixPlanError, match="no presentation brief"):
        build_master_plan(no_brief)


def test_source_identity_follows_the_approved_content_only() -> None:
    row = approved_row()
    manifest = generation_manifest(row)
    assert same_source(manifest, generation_manifest(copy.deepcopy(row)))
    edited = copy.deepcopy(row)
    edited["paper_json"]["analysis"]["research"]["core_thesis"]["text"] = "Working hypothesis: another thesis."
    edited["paper_json"]["presentation_brief"]["core_thesis"] = "Working hypothesis: another thesis."
    assert not same_source(manifest, generation_manifest(edited))
    assert not same_source(manifest, generation_manifest({**row, "id": "44444444-4444-4444-8444-444444444444"}))
    assert not same_source(manifest, {**manifest, "master_sha256": "0" * 64})
    assert not same_source(manifest, {"kind": "ppt1", "approved_discovery_version_id": row["id"]})
    # Page layout, generation state or a later draft are not part of the appendix source.
    relaid = copy.deepcopy(row)
    relaid["paper_json"]["page_manifest"] = []
    relaid["paper_json"]["latest_approved_version_id"] = "x"
    assert same_source(manifest, generation_manifest(relaid))


# --- appendix: structure ------------------------------------------------------------------


def test_plan_is_the_canonical_deck_followed_by_the_appendix() -> None:
    row = approved_row()
    plan = build_master_plan(row)
    slides = plan["slides"]
    assert (plan["engine"], plan["deck_kind"], plan["title"]) == ("borek_deck", "master_v1", "Master Presentation — Nordwind Maschinenbau")
    assert [slide["order"] for slide in slides] == list(range(1, len(slides) + 1))
    canonical, extra = slides[:26], slides[26:]
    assert {slide["layoutId"] for slide in canonical} == {CANONICAL_LAYOUT}
    assert [slide["borekSlide"]["master_slide"] for slide in canonical] == list(range(1, 27))
    assert all(slide["frameworkReferences"] == [] for slide in canonical)
    assert len(slides) > 8, "the old eight-slide cap does not apply"
    assert extra[0]["order"] == 27 and plan["appendix"] == {
        "first_slide": 27,
        "slide_count": len(extra),
        "source_hash": appendix.appendix_source_hash(row["paper_json"]),
    }
    assert all(slide["frameworkReferences"] for slide in extra)
    # A synthesis, not a slide per white-paper page or per opportunity.
    paper = row["paper_json"]
    opportunities = sum(len(area["opportunities"]) for area in paper["analysis"]["areas"])
    assert len(extra) < len(paper["page_manifest"]) and len(extra) < opportunities
    assert build_master_plan(approved_row()) == plan, "planning is deterministic"


# --- appendix: composition depends on the approved Discovery -------------------------------


def test_appendix_composition_follows_the_content_of_the_approved_discovery() -> None:
    minimal, medium, rich = (layouts_of(CASES[name]) for name in ("minimal", "medium", "rich"))
    assert minimal == ["L06", "L01", "L25", "L08", "L08", "L07"]
    assert medium == ["L06", "L01", "L25", "L08", "L08", "L14", "L08", "L07"]
    assert rich == ["L06", "L01", "L25", "L25", "L25", "L08", "L08", "L14", "L14", "L08", "L07"]
    assert len({len(minimal), len(medium), len(rich)}) == 3, "not every client gets the same number of slides"

    # Minimal: no client context, so no conversation-specific slides and nothing made up.
    slides = appendix.plan_appendix(paper_for(MINIMAL), first_order=27)
    purposes = [slide["purpose"] for slide in slides]
    assert "Open questions and decisions to validate" not in purposes and "Selected target workflows" not in purposes
    context = next(slide for slide in slides if slide["layout_id"] == "L01")["content"]
    assert context["left"]["bullets"] == ["No client information was provided for this analysis."]
    assert slides[0]["content"]["title"] == "Your AI opportunities"
    cards = next(slide for slide in slides if slide["layout_id"] == "L25")
    assert len({ref.split(":")[1].split("/")[0] for ref in cards["source_references"]}) == 4, "one opportunity per area"

    # Medium: the context points at one area, which is shown in depth and gets its questions.
    paper = paper_for(MEDIUM)
    slides = appendix.plan_appendix(paper, first_order=27)
    cards = next(slide for slide in slides if slide["layout_id"] == "L25")
    assert {ref.split(":")[1].split("/")[0] for ref in cards["source_references"]} == {"sales_quoting"}
    assert "areas your context points at" in cards["content"]["lead"]
    assert sum(1 for slide in slides if slide["layout_id"] == "L14") == 1

    # Rich: several focus areas add opportunity and question slides, numbered within their group.
    slides = appendix.plan_appendix(paper_for(RICH), first_order=27)
    titles = [slide["content"]["title"] for slide in slides if slide["layout_id"] == "L25"]
    assert titles == ["Where we would start (1/3)", "Where we would start (2/3)", "Where we would start (3/3)"]
    featured = [ref for slide in slides if slide["layout_id"] == "L25" for ref in slide["source_references"]]
    assert len(featured) == len(set(featured)) == 12, "no opportunity is repeated"

    # More areas in the analysis mean more area rows, never dropped ones.
    eight = paper_for(MEDIUM)
    extra = copy.deepcopy(eight["analysis"]["areas"][-1])
    extra["id"], extra["name"] = "field_service", "Field Service"
    for item in extra["opportunities"]:
        item["id"] = f"fs_{item['id']}"
    eight["analysis"]["areas"].append(extra)
    eight = finalize(eight)
    area_slides = [s for s in appendix.plan_appendix(eight, first_order=27) if s["purpose"] == "Priority business areas"]
    assert [len(slide["content"]["rows"]) for slide in area_slides] == [4, 4]
    assert [row["code"] for slide in area_slides for row in slide["content"]["rows"]] == [f"{n:02d}" for n in range(1, 9)]


@pytest.mark.parametrize("case", sorted(CASES))
def test_every_planned_slide_uses_an_implemented_layout_and_fits(case: str) -> None:
    paper = paper_for(CASES[case])
    slides = appendix.plan_appendix(paper, first_order=27)
    used = {slide["layout_id"] for slide in slides}
    assert used <= set(appendix.PLANNER_LAYOUTS) <= implemented_layouts() <= set(LAYOUT_IDS)
    assert assembly.fit_problems([slide["content"] for slide in slides]) == []
    assert [slide["order"] for slide in slides] == list(range(27, 27 + len(slides)))
    for slide in slides:
        assert set(slide) == {"order", "layout_id", "purpose", "content", "source_references"}
        for reference in slide["source_references"]:
            assert appendix.resolve_reference(paper, reference) is not None, reference
    rendered = json.dumps(slides, ensure_ascii=False)
    for forbidden in ("Confidential", "transcript", "personal notes", "Lorem", "Placeholder"):
        assert forbidden not in rendered


def test_planner_never_selects_a_layout_that_is_not_implemented(monkeypatch: pytest.MonkeyPatch) -> None:
    assert appendix.PLANNER_LAYOUTS == ("L06", "L01", "L25", "L08", "L14", "L07")
    assert implemented_layouts() == {"L01", "L06", "L07", "L08", "L14", "L25"}
    # If a planner layout lost its implementation, planning stops before any slide is produced.
    monkeypatch.setattr(appendix, "PLANNER_LAYOUTS", (*appendix.PLANNER_LAYOUTS, "L21"))
    with pytest.raises(appendix.AppendixPlanError, match="refers to a layout that is not implemented"):
        appendix.plan_appendix(paper_for(MEDIUM), first_order=27)


def test_appendix_content_comes_from_the_approved_analysis() -> None:
    paper = paper_for(MEDIUM)
    analysis = paper["analysis"]
    slides = appendix.plan_appendix(paper, first_order=27)
    by_purpose = {slide["purpose"]: slide for slide in slides}
    context = by_purpose["Client context and business thesis"]["content"]
    assert context["left"]["bullets"][0] == "Company name: Nordwind Maschinenbau"
    assert set(context["right"]["bullets"]) <= set(analysis["research"]["unknown_facts"])
    assert "Faster quoting for the sales team" in context["statement"]["text"]
    cards = by_purpose["Most important opportunities"]["content"]["cards"]
    titles = {item["title"] for area in analysis["areas"] for item in area["opportunities"]}
    assert len(cards) == 4 and {card["title"] for card in cards} <= titles
    asked = {q for area in analysis["areas"] for item in area["opportunities"] for q in item["discovery_questions"]}
    questions = by_purpose["Open questions and decisions to validate"]["content"]["items"]
    assert len(questions) == 4 and {item["question"] for item in questions} <= asked
    workflows = by_purpose["Selected target workflows"]["content"]["rows"]
    assert [row["title"] for row in workflows] == [item["name"] for item in analysis["target_workflows"]]
    steps = by_purpose["First pilot and validation approach"]["content"]["stages"]
    assert [stage["title"] for stage in steps] == ["Baseline workshop", "Pilot within weeks", "Scaling roadmap"]


@pytest.mark.parametrize(
    ("mutate", "message"),
    [
        (lambda s: s[1].__setitem__("layout_id", "L26"), "not a master layout"),
        (lambda s: s[1].__setitem__("layout_id", "COVER_01"), "not a master layout"),
        (lambda s: (s[1].__setitem__("layout_id", "L02"), s[1]["content"].__setitem__("layout", "L02")), "L02, which is not implemented yet"),
        (lambda s: s[2].__setitem__("source_references", []), "cites no source"),
        (lambda s: s[2].__setitem__("source_references", ["opportunity:other/thing"]), "not part of the approved analysis"),
        (lambda s: s[2].__setitem__("source_references", ["transcript_summary"]), "not part of the approved analysis"),
        (lambda s: s[1]["content"]["left"]["bullets"].__setitem__(0, "Revenue: 95 million"), "figures that are not in the approved analysis: 95"),
        (lambda s: s[2]["content"]["cards"][0].__setitem__("text", "Saves 1,300 hours per month."), "figures that are not in the approved analysis: 1,300"),
        (lambda s: s[3].__setitem__("order", 99), "has order 99"),
        (lambda s: s[2]["content"].__setitem__("title", "A title that is far too long to stay on one line in this layout, by a wide margin"), "title needs"),
        (lambda s: s.insert(len(s) - 1, {**copy.deepcopy(s[-1]), "order": s[-1]["order"]}) or s[-1].__setitem__("order", s[-1]["order"] + 1), "card-variety rule"),
        (lambda s: s[1].pop("purpose"), "lacks purpose"),
    ],
)
def test_invalid_appendix_slides_are_rejected(mutate, message: str) -> None:
    paper = paper_for(MEDIUM)
    slides = appendix.plan_appendix(paper, first_order=27)
    mutate(slides)
    with pytest.raises(appendix.AppendixPlanError, match=message):
        appendix.validate_appendix(slides, paper, first_order=27)


def test_layout_library_is_l01_to_l25_and_text_never_shrinks() -> None:
    assert LAYOUT_IDS == tuple(f"L{n:02d}" for n in range(1, 26))
    contents = [slide["content"] for slide in appendix.plan_appendix(paper_for(MEDIUM), first_order=27)]
    long_title = copy.deepcopy(contents)
    long_title[1]["title"] = "word " * 40
    assert any("title needs" in problem for problem in assembly.fit_problems(long_title))
    with pytest.raises(LayoutError, match="not implemented yet"):
        assembly.fit_problems([{"layout": "L21"}])
    with pytest.raises(LayoutError, match="not a layout of the master library"):
        assembly.fit_problems([{"layout": "cover"}])


# --- assembly -----------------------------------------------------------------------------


def test_assembly_appends_the_appendix_and_leaves_the_canonical_slides_untouched(tmp_path: Path) -> None:
    master = registry.get_master()
    contents = [slide["content"] for slide in appendix.plan_appendix(paper_for(RICH), first_order=27)]
    final = tmp_path / "final.pptx"
    assembly.assemble(master, contents, final)

    before = assembly.canonical_fingerprints(master.pptx_path)
    after = assembly.canonical_fingerprints(final)
    assert len(before) == 26 and len(after) == 26 + len(contents)
    assert after[:26] == before, "every canonical slide and its media are unchanged, in the original order"
    assert registry.file_sha256(master.pptx_path) == master.sha256, "the stored master itself is never written"

    source, deck = Presentation(str(master.pptx_path)), Presentation(str(final))
    text = lambda slide: [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]  # noqa: E731
    for number in range(26):
        assert text(deck.slides[number]) == text(source.slides[number])
    assert any("CONFIDENTIAL" in value.upper() for value in text(deck.slides[12])), "slide 13 keeps its footer"
    assert "01 / 26" in text(deck.slides[0]) and "26 / 26" in text(deck.slides[25]), "original footers are not renumbered"
    assert any("Managing Partner & Co-CCO" in value for value in text(deck.slides[25]))
    total = len(deck.slides)
    appendix_text = [value for number in range(26, total) for value in text(deck.slides[number])]
    assert f"27 / {total}" in appendix_text and f"{total} / {total}" in appendix_text
    assert not any("CONFIDENTIAL" in value.upper() for value in appendix_text)

    # No rounded corners, gradients, shadows or fonts other than Inter in the appendix.
    with zipfile.ZipFile(final) as package:
        parts = assembly._slide_parts(package)
        for part in parts[26:]:
            xml = package.read(part).decode("utf-8")
            assert "roundRect" not in xml and "gradFill" not in xml and "outerShdw" not in xml
            assert set(re.findall(r'typeface="([^"]+)"', xml)) <= {"Inter"}

    tampered = tmp_path / "tampered.pptx"
    with zipfile.ZipFile(final) as package, zipfile.ZipFile(tampered, "w", zipfile.ZIP_DEFLATED) as out:
        for item in package.infolist():
            data = package.read(item.filename)
            if item.filename == parts[5]:
                data = data.replace("Three ways in".encode(), "Four ways in".encode())
            out.writestr(item, data)
    with pytest.raises(assembly.AssemblyError, match="Canonical slide 6 was changed"):
        assembly.verify_canonical_preserved(master, tampered, expected_total=total)
    with pytest.raises(assembly.AssemblyError, match="expected 40"):
        assembly.verify_canonical_preserved(master, final, expected_total=40)


def _rewrite(source: Path, target: Path, part: str, change) -> None:
    """Copy a package, changing exactly one part; fails if the change did not apply."""
    with zipfile.ZipFile(source) as package, zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as out:
        assert part in package.namelist(), part
        for item in package.infolist():
            data = package.read(item.filename)
            if item.filename == part:
                changed = change(data)
                assert changed != data, f"the change did not apply to {part}"
                data = changed
            out.writestr(item, data)


def _sub(pattern: bytes, replacement: bytes):
    return lambda data: re.sub(pattern, replacement, data, count=1)


_MARK = _sub(rb"<p:cSld", b'<p:cSld xmlns:x="urn:changed" x:changed="1"')


def _dependency(deck: Path, start: str, kind: str) -> str:
    """Name of the part of the given kind that ``start`` references."""
    with zipfile.ZipFile(deck) as archive:
        targets = {target for _id, rel_kind, _external, target in assembly._Package(archive).relationships(start) if rel_kind == kind}
    return sorted(targets)[0]


@pytest.fixture(scope="module")
def assembled(tmp_path_factory: pytest.TempPathFactory) -> tuple[Path, int]:
    master = registry.get_master()
    contents = [slide["content"] for slide in appendix.plan_appendix(paper_for(MEDIUM), first_order=27)]
    final = tmp_path_factory.mktemp("assembled") / "final.pptx"
    assembly.assemble(master, contents, final)
    return final, 26 + len(contents)


def test_canonical_fingerprint_covers_the_whole_dependency_chain(assembled: tuple[Path, int]) -> None:
    final, total = assembled
    master = registry.get_master()
    with zipfile.ZipFile(master.pptx_path) as archive:
        package = assembly._Package(archive)
        closures = [{kind for kind, _checksum in package.closure(part)} for part in assembly._slide_parts(archive)]
    for kinds in closures:
        assert {"slide", "slideLayout", "slideMaster", "theme"} <= kinds
    assert {"notesSlide", "notesMaster"} <= closures[0]
    assert "image" in set().union(*closures)
    settings = assembly.presentation_settings(master.pptx_path, slide_count=26)
    assert {"presentation.xml", "slideMaster", "theme", "notesMaster", "presProps", "viewProps", "tableStyles"} <= {
        key.split(" ")[0] for key in settings
    }
    assert sum(key.startswith("slide ") for key in settings) == 26
    # Appending slides changes neither the canonical slides nor the presentation settings.
    assert assembly.presentation_settings(final, slide_count=26) == settings
    assert assembly.canonical_fingerprints(final)[:26] == assembly.canonical_fingerprints(master.pptx_path)
    assembly.verify_canonical_preserved(master, final, expected_total=total)


@pytest.mark.parametrize(
    ("case", "message"),
    [
        ("theme", r"Canonical slide 1 was changed during assembly \(its theme\)"),
        ("theme_content_type", r"Canonical slide 1 was changed during assembly \(its theme\)"),
        ("layout", r"Canonical slide 1 was changed during assembly \(its slideLayout\)"),
        ("layout_relationships", r"Canonical slide 1 was changed during assembly \(its slideLayout\)"),
        ("master", r"Canonical slide 1 was changed during assembly \(its slideMaster\)"),
        ("notes", r"Canonical slide 1 was changed during assembly \(its notesSlide\)"),
        ("media", r"Canonical slide \d+ was changed during assembly \(its image\)"),
        ("slide_size", r"presentation settings of the canonical deck were changed during assembly \(presentation.xml\)"),
        ("table_styles", r"presentation settings of the canonical deck were changed during assembly \(tableStyles\)"),
        ("slide_order", r"Canonical slide 1 was changed during assembly"),
    ],
)
def test_a_change_to_anything_a_canonical_slide_depends_on_is_detected(
    assembled: tuple[Path, int], tmp_path: Path, case: str, message: str
) -> None:
    final, total = assembled
    master = registry.get_master()
    with zipfile.ZipFile(final) as archive:
        slides = assembly._slide_parts(archive)
    layout = _dependency(final, slides[0], "slideLayout")
    slide_master = _dependency(final, layout, "slideMaster")
    external = b'<Relationship Id="rId99" Type="http://x/hyperlink" Target="http://x" TargetMode="External"/></Relationships>'
    cases = {
        "theme": (_dependency(final, slide_master, "theme"), _sub(rb'(<a:srgbClr val=")[0-9A-Fa-f]{6}', rb"\g<1>123456")),
        "theme_content_type": (
            "[Content_Types].xml",
            _sub(rb"application/vnd\.openxmlformats-officedocument\.theme\+xml", b"application/xml"),
        ),
        "layout": (layout, _MARK),
        "layout_relationships": (assembly._rels_part(layout), _sub(rb"</Relationships>", external)),
        "master": (slide_master, _MARK),
        "notes": (_dependency(final, slides[0], "notesSlide"), _MARK),
        "media": (_dependency(final, slides[2], "image"), lambda data: data + b"\0"),
        "slide_size": ("ppt/presentation.xml", _sub(rb'(<p:sldSz cx=")\d+', rb"\g<1>9144000")),
        "table_styles": (
            _dependency(final, "ppt/presentation.xml", "tableStyles"),
            _sub(rb"<a:tblStyleLst", b'<a:tblStyleLst xmlns:x="urn:changed" x:changed="1"'),
        ),
        "slide_order": ("ppt/presentation.xml", _sub(rb"(<p:sldId [^>]*/>)(<p:sldId [^>]*/>)", rb"\g<2>\g<1>")),
    }
    part, change = cases[case]
    tampered = tmp_path / "tampered.pptx"
    _rewrite(final, tampered, part, change)
    with pytest.raises(assembly.AssemblyError, match=message):
        assembly.verify_canonical_preserved(master, tampered, expected_total=total)


def test_assembly_refuses_to_write_a_deck_whose_slide_master_was_changed(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """A build that restyles the shared slide master changes all 26 slides: it must fail."""
    master = registry.get_master()
    contents = [slide["content"] for slide in appendix.plan_appendix(paper_for(MINIMAL), first_order=27)]
    original = assembly._draw_all

    def draw_and_restyle_master(deck: Any, layout: Any, *args: Any, **kwargs: Any) -> list[str]:
        problems = original(deck, layout, *args, **kwargs)
        deck.slide_masters[0].background.fill.solid()
        deck.slide_masters[0].background.fill.fore_color.rgb = assembly.RGBColor(0xFF, 0x00, 0x00)
        return problems

    monkeypatch.setattr(assembly, "_draw_all", draw_and_restyle_master)
    with pytest.raises(assembly.AssemblyError, match=r"Canonical slide 1 was changed during assembly \(its slideMaster\)"):
        assembly.assemble(master, contents, tmp_path / "final.pptx")
    assert registry.file_sha256(master.pptx_path) == master.sha256


# --- rendering ----------------------------------------------------------------------------


def test_render_bundle_is_built_from_the_assembled_deck() -> None:
    master = registry.get_master()
    specs = specs_of(build_master_plan(approved_row()))
    bundle = zipfile.ZipFile(io.BytesIO(render_deck_bundle(specs, deck_kind=MASTER_V1)))
    manifest = json.loads(bundle.read("manifest.json"))
    total = len(specs)
    assert manifest["slideCount"] == total == 26 + manifest["appendixSlideCount"]
    assert (manifest["pdfEngine"], manifest["previewEngine"], manifest["renderSource"]) == (
        "libreoffice", "libreoffice+pdftoppm", "deck.pptx"
    )
    assert manifest["master"] == {**master.identity(), "slideCount": 26, "canonicalSlidesVerified": True}
    assert manifest["previews"] == [f"preview-{n:03d}.png" for n in range(1, total + 1)]
    assert sorted(bundle.namelist()) == sorted(["deck.pptx", "deck.pdf", "manifest.json", *manifest["previews"]])
    assert len(Presentation(io.BytesIO(bundle.read("deck.pptx"))).slides) == total
    pdf = bundle.read("deck.pdf")
    assert pdf.startswith(b"%PDF") and len(re.findall(rb"/Type\s*/Page[^s]", pdf)) == total, "one PDF page per slide"


def test_failed_or_incomplete_conversion_fails_the_render_and_never_falls_back(monkeypatch: pytest.MonkeyPatch) -> None:
    specs = specs_of(build_master_plan(approved_row()))

    def unavailable(*_args, **_kwargs):
        raise office.OfficeRendererUnavailable("LibreOffice is not installed in this runtime; the deck cannot be rendered")

    monkeypatch.setattr(office, "render_pptx", unavailable)
    with pytest.raises(BorekDeckRenderError, match="LibreOffice is not installed") as failure:
        render_deck_bundle(specs, deck_kind=MASTER_V1)
    assert failure.value.code == "MASTER_RENDERER_UNAVAILABLE"

    def broken(*_args, **_kwargs):
        raise office.OfficeRenderError("The PDF has 30 pages for 34 slides")

    monkeypatch.setattr(office, "render_pptx", broken)
    with pytest.raises(BorekDeckRenderError, match="30 pages for 34 slides") as failure:
        render_deck_bundle(specs, deck_kind=MASTER_V1)
    assert failure.value.code == "MASTER_PRESENTATION_RENDER_FAILED"

    # The approximate Pillow renderer of the legacy decks is not involved in a Master Presentation.
    source = Path(render_master_bundle.__code__.co_filename).read_text(encoding="utf-8")
    assert "preview.render" not in source and "_pdf_from_previews" not in source and "PIL" not in source


def test_office_renderer_reports_missing_tools_explicitly(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    monkeypatch.setattr(office.shutil, "which", lambda _name: None)
    assert office.renderer_available() is False
    with pytest.raises(office.OfficeRendererUnavailable, match="LibreOffice is not installed"):
        REAL_RENDER_PPTX(tmp_path / "deck.pptx", tmp_path, expected_slides=1)


def test_render_rejects_decks_that_are_not_a_complete_master_presentation() -> None:
    specs = specs_of(build_master_plan(approved_row()))
    reordered = [specs[1], specs[0], *specs[2:]]
    with pytest.raises(BorekDeckRenderError, match="canonical slides in their original order"):
        render_deck_bundle(reordered, deck_kind=MASTER_V1)
    with pytest.raises(BorekDeckRenderError, match="canonical slides in their original order"):
        render_deck_bundle(specs[1:], deck_kind=MASTER_V1)
    with pytest.raises(BorekDeckRenderError, match="no canonical master slides"):
        render_deck_bundle(specs[26:], deck_kind=MASTER_V1)
    foreign = copy.deepcopy(specs)
    foreign[3]["content"]["master_id"] = "another_master"
    with pytest.raises(BorekDeckRenderError, match="different master decks|Unknown master deck"):
        render_deck_bundle(foreign, deck_kind=MASTER_V1)


# --- the real renderer (runs where LibreOffice, poppler and Inter are installed: the images) ---

real_renderer = pytest.mark.skipif(not office.renderer_available(), reason="LibreOffice/poppler are not installed here")


@real_renderer
@pytest.mark.parametrize("case", sorted(CASES))
def test_libreoffice_renders_the_assembled_deck_in_inter(case: str, tmp_path: Path) -> None:
    from PIL import Image

    master = registry.get_master()
    contents = [slide["content"] for slide in appendix.plan_appendix(paper_for(CASES[case]), first_order=27)]
    deck = tmp_path / "deck.pptx"
    assembly.assemble(master, contents, deck)
    total = 26 + len(contents)
    pdf, previews = office.render_pptx(deck, tmp_path, expected_slides=total)
    assert office.pdf_page_count(pdf) == total == len(previews)
    fonts = office.pdf_fonts(pdf)
    assert fonts and all(name.lower().startswith("inter") for name in fonts), fonts
    for path in previews:
        with Image.open(path) as image:
            assert image.size == (1920, 1080)
            assert len(image.convert("L").getcolors(maxcolors=1 << 20) or [0, 0]) > 1, f"{path.name} is blank"
    with pytest.raises(office.OfficeRenderError, match="pages for"):
        office.render_pptx(deck, tmp_path / "again", expected_slides=total + 1)


@real_renderer
def test_a_deck_that_is_not_a_presentation_fails_conversion(tmp_path: Path) -> None:
    broken = tmp_path / "deck.pptx"
    broken.write_bytes(b"this is not a PowerPoint package")
    with pytest.raises(office.OfficeRenderError):
        office.render_pptx(broken, tmp_path, expected_slides=1)


@real_renderer
@pytest.mark.parametrize("case", sorted(CASES))
def test_no_text_overlaps_on_any_rendered_appendix_slide(case: str, tmp_path: Path) -> None:
    """Measured in the real render: two text lines of an appendix slide never share space."""
    import subprocess
    from xml.etree import ElementTree

    master = registry.get_master()
    contents = [slide["content"] for slide in appendix.plan_appendix(paper_for(CASES[case]), first_order=27)]
    deck = tmp_path / "deck.pptx"
    assembly.assemble(master, contents, deck)
    pdf, _previews = office.render_pptx(deck, tmp_path, expected_slides=26 + len(contents))
    layout = tmp_path / "layout.html"
    subprocess.run(["pdftotext", "-bbox-layout", "-f", "27", str(pdf), str(layout)], check=True, timeout=120)
    root = ElementTree.parse(layout).getroot()
    pages = [node for node in root.iter() if node.tag.endswith("page")]
    assert len(pages) == len(contents)
    for number, page in enumerate(pages, start=27):
        lines = [
            (float(n.get("xMin")), float(n.get("yMin")), float(n.get("xMax")), float(n.get("yMax")), " ".join(w.text or "" for w in n))
            for n in page.iter()
            if n.tag.endswith("line")
        ]
        assert lines, f"slide {number} has no text"
        height = float(page.get("height"))
        assert all(y1 <= height for _x0, _y0, _x1, y1, _t in lines), f"text runs off slide {number}"
        for index, (ax0, ay0, ax1, ay1, text_a) in enumerate(lines):
            for bx0, by0, bx1, by1, text_b in lines[index + 1 :]:
                overlap_x = min(ax1, bx1) - max(ax0, bx0)
                overlap_y = min(ay1, by1) - max(ay0, by0)
                smaller = min(ay1 - ay0, by1 - by0)
                assert not (overlap_x > 2 and overlap_y > 0.35 * smaller), (
                    f"slide {number}: '{text_a[:40]}' overlaps '{text_b[:40]}'"
                )
