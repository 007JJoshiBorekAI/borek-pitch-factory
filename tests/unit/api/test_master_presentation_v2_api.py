"""Master Presentation V2 through the API: identity, frozen sources, idempotency, review, finalization."""

from __future__ import annotations

import copy
import io
import json
from uuid import UUID

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

from app.auth import create_test_access_token
from app.config import settings
from app.main import create_app
from app.services import presentation_generation
from app.services.data.memory_store import get_memory_store, reset_memory_store
from services.presentation.master_deck import assembly, office, registry

OWNER = UUID("aaaaaaaa-aaaa-4aaa-8aaa-aaaaaaaaaaaa")
OTHER = UUID("bbbbbbbb-bbbb-4bbb-8bbb-bbbbbbbbbbbb")
SUPPLIED_SHA256 = "2c23670e46dc2ce71d3527b0acc846cad54572d9817977ec26c3efe17f51dd5a"
CLIENT = {
    "client_name": "Nordwind Maschinenbau",
    "opportunity_name": "AI introduction",
    "department": "Sales",
    "stage1_intake": {
        "poc_name": "Dana Weber",
        "sales_topic_description": "Faster quoting for the sales team",
        "about_company": "Family-owned machine builder with 420 employees.",
    },
}
MEETING = (
    "Dana: Requirement: Quotes must go out within one day.\n"
    "Tom: Challenge: Pricing data sits in three systems.\n"
    "Dana: Priority: Start with the sales team.\n"
    "Tom: Discussed solution: A drafting assistant that prepares the quote from the request.\n"
    "Tom: Decision: Run a pilot with ten quotes.\n"
    "Dana: Follow-up: Send the pricing export by Friday.\n"
    "Tom: We also talked about the trade fair.\n"
)
NOTES = "Priority: Start with the sales team.\nOpportunity: The service team could reuse the quoting assistant."
EXCLUDED = "Send the pricing export by Friday."
OBSERVATION = "The service team could reuse the quoting assistant."


def headers(user: UUID = OWNER) -> dict[str, str]:
    token = create_test_access_token(user_id=user, email="sales@example.com", secret=settings.SUPABASE_JWT_SECRET)
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture()
def client() -> TestClient:
    reset_memory_store()
    return TestClient(create_app())


def ok(response, status: int = 200) -> dict:
    assert response.status_code == status, response.text
    return response.json()


def get(client: TestClient, path: str, user: UUID = OWNER):
    return client.get(path, headers=headers(user))


def post(client: TestClient, path: str, body: dict | None = None, user: UUID = OWNER):
    return client.post(path, headers=headers(user), json=body)


def review(client: TestClient, opportunity_id: str) -> dict:
    return ok(get(client, f"/opportunities/{opportunity_id}/post-meeting-review"))


def confirm(client: TestClient, opportunity_id: str, excluded: dict | None = None) -> dict:
    view = review(client, opportunity_id)
    return ok(
        post(
            client,
            f"/opportunities/{opportunity_id}/post-meeting-review/confirm",
            {
                "transcript_id": view["extraction"]["transcript_id"],
                "extraction_generated_at": view["extraction"]["generated_at"],
                "review_fingerprint": view["review_fingerprint"],
                "excluded": excluded or {},
            },
        )
    )


def analyse(client: TestClient, opportunity_id: str, transcript_id: str) -> None:
    ok(post(client, f"/opportunities/{opportunity_id}/meeting-extraction/generate", {"transcript_id": transcript_id}))


def prepared(client: TestClient, user: UUID = OWNER, *, body: dict = CLIENT, meeting: str = MEETING, notes: str | None = NOTES) -> dict:
    """Approved Discovery v2, ready Master Presentation V1, first meeting, analysed and confirmed findings."""
    opportunity_id = ok(client.post("/opportunities", headers=headers(user), json=body), 201)["id"]
    ok(post(client, f"/opportunities/{opportunity_id}/discovery-paper/generate", user=user))
    approved = ok(post(client, f"/opportunities/{opportunity_id}/discovery-paper/approve", user=user))
    deck = ok(post(client, f"/opportunities/{opportunity_id}/stage1-outputs/generate", user=user))["outputs"]["presentation"]
    assert deck["status"] == "ready", deck
    ok(post(client, f"/opportunities/{opportunity_id}/workflow/first-meeting-completed", user=user))
    uploaded = client.post(
        f"/opportunities/{opportunity_id}/transcripts",
        headers=headers(user),
        files={"file": ("meeting.txt", io.BytesIO(meeting.encode()), "text/plain")},
    )
    transcript_id = ok(uploaded, 201)["transcript"]["id"]
    if notes is not None:
        ok(client.put(f"/opportunities/{opportunity_id}/personal-notes", headers=headers(user), json={"text": notes}))
    ok(post(client, f"/opportunities/{opportunity_id}/meeting-extraction/generate", {"transcript_id": transcript_id}, user))
    return {"opportunity": opportunity_id, "presentation": deck["presentation_id"], "approved": approved["id"], "transcript": transcript_id}


def generate_v2(client: TestClient, opportunity_id: str, user: UUID = OWNER):
    return post(client, f"/opportunities/{opportunity_id}/master-presentation/v2/generate", user=user)


def v2_status(client: TestClient, opportunity_id: str) -> dict:
    return ok(get(client, f"/opportunities/{opportunity_id}/master-presentation/v2"))


def versions(client: TestClient, presentation_id: str) -> list[dict]:
    return ok(get(client, f"/presentations/{presentation_id}/versions"))


def stored_versions(presentation_id: str) -> list[dict]:
    rows = [row for row in get_memory_store().presentation_versions.values() if str(row["presentation_id"]) == presentation_id]
    return sorted(rows, key=lambda row: row["version_number"])


def slide_texts(pptx: bytes) -> list[list[str]]:
    deck = Presentation(io.BytesIO(pptx))
    return [[shape.text_frame.text for shape in slide.shapes if shape.has_text_frame] for slide in deck.slides]


def download(client: TestClient, presentation_id: str, version_id: str, kind: str = "pptx") -> bytes:
    response = get(client, f"/presentations/{presentation_id}/versions/{version_id}/download/{kind}")
    assert response.status_code == 200, response.text
    return response.content


def fingerprints(pptx: bytes, tmp_path) -> list[str]:
    path = tmp_path / "deck.pptx"
    path.write_bytes(pptx)
    return assembly.canonical_fingerprints(path)


def test_v2_is_a_new_version_of_the_same_presentation_with_a_post_meeting_appendix(client: TestClient, tmp_path) -> None:
    base = prepared(client)
    opportunity_id, presentation_id = base["opportunity"], base["presentation"]
    (v1_row,) = stored_versions(presentation_id)
    v1_snapshot = copy.deepcopy(v1_row)
    v1_files = {kind: download(client, presentation_id, str(v1_row["id"]), kind) for kind in ("pptx", "pdf")}
    ready = confirm(client, opportunity_id, {"follow_ups": [EXCLUDED]})
    assert ready["readiness"]["ready_for_v2"]
    assert v2_status(client, opportunity_id)["state"] == "none"

    started = ok(generate_v2(client, opportunity_id))
    assert (started["presentation_id"], started["product_version"], started["status"]) == (presentation_id, "V2", "COMPLETED")
    assert started["base_presentation_version_id"] == str(v1_row["id"])
    assert started["source_hash"] == ready["v2_sources"]["source_hash"]

    # One presentation, two versions. The product version is not the revision number.
    store = get_memory_store()
    assert len(store.presentations) == 1
    v1_after, v2_row = stored_versions(presentation_id)
    assert v1_after == v1_snapshot, "the V1 version and its manifest are untouched"
    assert (v2_row["version_number"], v2_row["journey_stage"], v2_row["status"]) == (2, "post_meeting", "ready")
    assert str(v2_row["prior_stage_presentation_version_id"]) == str(v1_row["id"])
    manifest = v2_row["generation_source_manifest"]
    assert (manifest["kind"], manifest["product_version"], manifest["product_stage"]) == ("master_presentation_v2", "V2", "post_meeting")
    assert manifest["master_sha256"] == SUPPLIED_SHA256
    assert manifest["presentation_id"] == presentation_id and manifest["base_presentation_version_id"] == str(v1_row["id"])
    assert manifest["approved_discovery_version_id"] == base["approved"]
    assert manifest["transcript_id"] == base["transcript"] and manifest["transcript_revision"] == ready["v2_sources"]["transcript_revision"]
    assert manifest["review_fingerprint"] == ready["review_fingerprint"]
    assert (manifest["confirmed_finding_count"], manifest["excluded_finding_count"]) == (6, 1)
    assert "Quotes must" not in json.dumps(manifest), "the manifest carries identities and checksums, not content"

    listed = versions(client, presentation_id)
    assert [(item["version_number"], item["is_latest"], item["source"]["product_version"]) for item in listed] == [(2, True, "V2"), (1, False, "V1")]
    assert listed[0]["source"]["base_presentation_version_id"] == str(v1_row["id"])
    assert "base_presentation_version_id" not in listed[1]["source"]

    # V1 files are byte-identical; V2 is a complete deck: the canonical 26 slides and a new appendix.
    for kind in ("pptx", "pdf"):
        assert download(client, presentation_id, str(v1_row["id"]), kind) == v1_files[kind]
    v2_pptx = download(client, presentation_id, str(v2_row["id"]))
    canonical = assembly.canonical_fingerprints(registry.get_master().pptx_path)
    v1_prints, v2_prints = fingerprints(v1_files["pptx"], tmp_path), fingerprints(v2_pptx, tmp_path)
    assert v1_prints[:26] == canonical and v2_prints[:26] == canonical
    assert v2_prints[26:] != v1_prints[26:], "the appendix is new, not the V1 appendix again"
    v1_texts, v2_texts = slide_texts(v1_files["pptx"]), slide_texts(v2_pptx)
    assert v2_texts[:26] == v1_texts[:26]
    total = len(v2_texts)
    assert f"27 / {total}" in v2_texts[26] and f"{total} / {total}" in v2_texts[-1]
    appendix = v2_texts[26:]
    flat = [text for slide in appendix for text in slide]
    assert any("AFTER OUR FIRST MEETING" in text.upper() for text in appendix[0])
    for said in ("Quotes must go out within one day.", "Pricing data sits in three systems.", "Run a pilot with ten quotes."):
        assert said in flat, said
    assert EXCLUDED not in " ".join(text for slide in v2_texts for text in slide), "an excluded finding is never used"
    assert "trade fair" not in " ".join(flat)
    # The owner's observation appears once, on its own page, never among the client statements.
    holders = [slide for slide in appendix if OBSERVATION in slide]
    assert len(holders) == 1 and any("NOT CLIENT STATEMENTS" in text.upper() for text in holders[0])
    client_pages = [slide for slide in appendix if any(text.upper().startswith("FROM THE MEETING") for text in slide)]
    assert client_pages and all(OBSERVATION not in slide for slide in client_pages)

    # Deck Center: latest is V2; each version is served with its own identity and slides.
    latest = ok(get(client, f"/presentations/{presentation_id}/deck"))
    assert (latest["version_number"], latest["source"]["product_version"], latest["source"]["product_stage"]) == (2, "V2", "post_meeting")
    assert latest["source"]["canonical_slide_count"] == 26 and latest["source"]["appendix_slide_count"] == total - 26
    for row, product in ((v1_row, "V1"), (v2_row, "V2")):
        prefix = f"/presentations/{presentation_id}/versions/{row['id']}"
        deck = ok(get(client, f"{prefix}/deck"))
        slides = ok(get(client, f"{prefix}/slides"))
        assert deck["source"]["product_version"] == product
        assert {slide["presentation_version_id"] for slide in slides} == {str(row["id"])}
        assert [slide["id"] for slide in slides] == [slide["slide_id"] for slide in deck["slides"]]
        assert get(client, f"{prefix}/preview/slides/26.png").status_code == 200
        assert get(client, f"{prefix}/slides", OTHER).status_code in (403, 404)

    # Post Meeting still finds V1, and its confirmation is not made stale by its own V2.
    after = review(client, opportunity_id)
    assert after["master_presentation"] == ready["master_presentation"]
    assert after["master_presentation"]["version_id"] == str(v1_row["id"])
    assert (after["confirmation"]["status"], after["readiness"]["ready_for_v2"]) == ("current", True)
    assert after["review_fingerprint"] == ready["review_fingerprint"] and after["v2_sources"] == ready["v2_sources"]
    workflow = ok(get(client, f"/opportunities/{opportunity_id}/workflow-status"))
    assert workflow["documents"]["ppt1"] == {
        "presentation_id": presentation_id, "latest_ready_version_id": str(v1_row["id"]),
        "journey_stage": "first_contact", "status": "ready", "product_version": "V1",
    }
    assert workflow["documents"]["ppt2"] == {
        "presentation_id": presentation_id, "latest_ready_version_id": str(v2_row["id"]),
        "journey_stage": "post_meeting", "status": "ready", "product_version": "V2",
    }
    assert v2_status(client, opportunity_id)["state"] == "ready"
    # No standalone PPT #2 was involved, and the Stage 1 pointer still names the same presentation.
    assert all(_job_kind(row) != "ppt2" for row in store.generation_jobs.values())
    assert ok(get(client, f"/opportunities/{opportunity_id}/stage1-outputs"))["outputs"]["presentation"]["presentation_id"] == presentation_id


def _job_kind(row: dict) -> str | None:
    manifest = ((row.get("result_json") or {}).get("_enqueue") or {}).get("generation_source_manifest") or {}
    return manifest.get("kind")


def test_generation_is_refused_until_the_backend_reports_the_review_as_ready(client: TestClient) -> None:
    base = prepared(client)
    opportunity_id, presentation_id = base["opportunity"], base["presentation"]
    before = copy.deepcopy(stored_versions(presentation_id))
    jobs_before = len(get_memory_store().generation_jobs)

    def refused(expected: str) -> None:
        response = generate_v2(client, opportunity_id)
        assert response.status_code == 409 and "MASTER_V2_NOT_READY" in response.text and expected in response.text, response.text
        assert stored_versions(presentation_id) == before and len(get_memory_store().generation_jobs) == jobs_before

    refused("confirm the meeting findings")  # analysed, but nothing confirmed
    assert generate_v2(client, opportunity_id, OTHER).status_code in (403, 404)
    assert client.post(f"/opportunities/{opportunity_id}/master-presentation/v2/generate").status_code == 401
    assert client.get(f"/opportunities/{opportunity_id}/master-presentation/v2").status_code == 401
    assert get(client, f"/opportunities/{opportunity_id}/master-presentation/v2", OTHER).status_code in (403, 404)
    status = v2_status(client, opportunity_id)
    assert (status["state"], status["can_generate"], status["blockers"]) == ("none", False, ["MEETING_REVIEW_NOT_CONFIRMED"])

    confirm(client, opportunity_id)
    ok(client.put(f"/opportunities/{opportunity_id}/personal-notes", headers=headers(), json={"text": NOTES + "\nDecision: Owner sponsors the pilot."}))
    refused("analyse the meeting again")  # confirmed, then a source changed
    analyse(client, opportunity_id, base["transcript"])
    refused("confirm the meeting findings again")  # analysed again, confirmation is for the old analysis
    confirm(client, opportunity_id)
    assert v2_status(client, opportunity_id)["can_generate"] is True
    assert ok(generate_v2(client, opportunity_id))["status"] == "COMPLETED"

    # A review with nothing confirmed cannot become a presentation.
    second = prepared(client)
    everything = {
        category: [item["text"] for item in items]
        for category, items in review(client, second["opportunity"])["extraction"]["categories"].items()
    }
    view = review(client, second["opportunity"])
    nothing = post(
        client,
        f"/opportunities/{second['opportunity']}/post-meeting-review/confirm",
        {
            "transcript_id": view["extraction"]["transcript_id"],
            "extraction_generated_at": view["extraction"]["generated_at"],
            "review_fingerprint": view["review_fingerprint"],
            "excluded": everything,
        },
    )
    assert nothing.status_code == 400 and "MEETING_REVIEW_NO_FINDINGS" in nothing.text, nothing.text
    # The refused confirmation changed nothing: the earlier one still stands.
    assert review(client, second["opportunity"])["confirmation"] == view["confirmation"]
    assert len(stored_versions(second["presentation"])) == 1


def test_same_sources_return_the_same_version_and_new_confirmation_adds_a_revision(client: TestClient) -> None:
    base = prepared(client)
    opportunity_id, presentation_id = base["opportunity"], base["presentation"]
    confirm(client, opportunity_id, {"follow_ups": [EXCLUDED]})
    first = ok(generate_v2(client, opportunity_id))
    store = get_memory_store()
    snapshot = copy.deepcopy(stored_versions(presentation_id))
    jobs = len(store.generation_jobs)

    for _ in range(3):
        again = ok(generate_v2(client, opportunity_id))
        assert (again["presentation_version_id"], again["status"], again["job_id"]) == (first["presentation_version_id"], "ready", None)
        assert again["is_existing_version"] is True
    assert stored_versions(presentation_id) == snapshot and len(store.generation_jobs) == jobs
    assert v2_status(client, opportunity_id)["can_generate"] is False, "nothing to generate: V2 matches the confirmation"

    # Confirming the same findings with the same exclusions again is recorded, but it is no new input:
    # the ready V2 stays current and no version is generated.
    before = review(client, opportunity_id)["confirmation"]["confirmed_at"]
    reconfirmed = confirm(client, opportunity_id, {"follow_ups": [EXCLUDED]})
    assert reconfirmed["confirmation"]["confirmed_at"] != before, "the new confirmation is kept for traceability"
    assert len([e for e in store.audit_logs.values() if e["action"] == "meeting_review.confirm"]) == 2
    status = v2_status(client, opportunity_id)
    assert (status["state"], status["can_generate"]) == ("ready", False)
    assert status["latest_ready"]["presentation_version_id"] == first["presentation_version_id"]
    assert status["latest_ready"]["generation_fingerprint"] == status["current_generation_fingerprint"]
    again = ok(generate_v2(client, opportunity_id))
    assert (again["presentation_version_id"], again["is_existing_version"]) == (first["presentation_version_id"], True)
    # The same holds after the meeting was analysed again with unchanged sources.
    analyse(client, opportunity_id, base["transcript"])
    confirm(client, opportunity_id, {"follow_ups": [EXCLUDED]})
    assert v2_status(client, opportunity_id)["state"] == "ready"
    assert ok(generate_v2(client, opportunity_id))["presentation_version_id"] == first["presentation_version_id"]
    assert stored_versions(presentation_id) == snapshot and len(store.generation_jobs) == jobs
    # The stored manifest keeps the confirmation it was generated with; it is never rewritten.
    assert snapshot[1]["generation_source_manifest"]["meeting_review_confirmed_at"] == before

    # Newly confirmed information: the ready V2 is reported as outdated until a new revision is requested.
    confirm(client, opportunity_id)  # the follow-up is included now
    status = v2_status(client, opportunity_id)
    assert (status["state"], status["can_generate"]) == ("outdated", True)
    assert status["latest_ready"]["presentation_version_id"] == first["presentation_version_id"]
    second = ok(generate_v2(client, opportunity_id))
    assert second["presentation_id"] == presentation_id and second["presentation_version_id"] != first["presentation_version_id"]
    v1_row, v2_first, v2_second = stored_versions(presentation_id)
    assert [v1_row, v2_first] == snapshot, "earlier versions are never rewritten"
    assert (v2_second["version_number"], v2_second["generation_source_manifest"]["product_version"]) == (3, "V2")
    assert v2_second["generation_source_manifest"]["base_presentation_version_id"] == str(v1_row["id"])
    assert v2_second["generation_source_manifest"]["excluded_finding_count"] == 0
    listed = versions(client, presentation_id)
    assert [(item["version_number"], item["source"]["product_version"], item["source"]["revision"]) for item in listed] == [
        (3, "V2", 3), (2, "V2", 2), (1, "V1", 1),
    ]
    first_texts = " ".join(text for slide in slide_texts(download(client, presentation_id, first["presentation_version_id"])) for text in slide)
    second_texts = " ".join(text for slide in slide_texts(download(client, presentation_id, second["presentation_version_id"])) for text in slide)
    assert EXCLUDED not in first_texts and EXCLUDED in second_texts
    assert len(store.presentations) == 1 and v2_status(client, opportunity_id)["state"] == "ready"


def test_a_running_job_is_reused_and_the_worker_renders_the_frozen_sources(client: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    base = prepared(client)
    opportunity_id, presentation_id = base["opportunity"], base["presentation"]
    confirm(client, opportunity_id, {"follow_ups": [EXCLUDED]})
    store = get_memory_store()
    dispatched: list[tuple] = []
    dispatch_now = presentation_generation._dispatch_task
    monkeypatch.setattr(presentation_generation, "_dispatch_task", lambda task, *args: dispatched.append((task, args)))

    queued = ok(generate_v2(client, opportunity_id))
    assert (queued["status"], queued["presentation_version_id"], queued["is_existing_job"]) == ("QUEUED", None, False)
    assert v2_status(client, opportunity_id)["state"] == "generating"
    # Double click / reload: the same job, nothing new.
    again = ok(generate_v2(client, opportunity_id))
    assert (again["job_id"], again["is_existing_job"]) == (queued["job_id"], True)
    assert len(dispatched) == 1 and len(stored_versions(presentation_id)) == 1
    job_row = store.generation_jobs[UUID(queued["job_id"])]
    frozen = job_row["result_json"]["_enqueue"]["master_v2_snapshot"]
    assert frozen["approved_discovery"]["version_id"] == base["approved"]
    assert frozen["personal_notes"]["text"] == NOTES
    assert frozen["transcript"]["transcript_id"] == base["transcript"]
    assert frozen["excluded_findings"] == [EXCLUDED]
    assert [item["text"] for item in frozen["findings"]["follow_ups"]] == []
    assert {"text": OBSERVATION, "source": "personal_notes"} in frozen["findings"]["opportunities"]

    # While the job waits, every live source changes. A different V2 cannot start meanwhile.
    ok(client.put(f"/opportunities/{opportunity_id}/personal-notes", headers=headers(), json={"text": "Decision: Cancel everything."}))
    analyse(client, opportunity_id, base["transcript"])
    edited = client.patch(
        f"/opportunities/{opportunity_id}/discovery-paper",
        headers=headers(),
        json={"edits": [{"target": "thesis", "value": {"text": "Working hypothesis: approved after the request"}}]},
    )
    assert edited.status_code == 200, edited.text
    newer = ok(post(client, f"/opportunities/{opportunity_id}/discovery-paper/approve"))
    confirm(client, opportunity_id)
    busy = generate_v2(client, opportunity_id)
    assert busy.status_code == 409 and "PRESENTATION_GENERATION_IN_PROGRESS" in busy.text

    # The worker runs now and builds exactly what was frozen at the request.
    task, args = dispatched[0]
    task(*args)
    v1_row, v2_row = stored_versions(presentation_id)
    manifest = v2_row["generation_source_manifest"]
    assert v2_row["status"] == "ready"
    assert manifest["approved_discovery_version_id"] == base["approved"] != newer["id"]
    assert manifest["snapshot_hash"] == job_row["result_json"]["_enqueue"]["generation_source_manifest"]["snapshot_hash"]
    text = " ".join(t for slide in slide_texts(download(client, presentation_id, str(v2_row["id"]))) for t in slide)
    assert "Cancel everything" not in text and "approved after the request" not in text
    assert "Quotes must go out within one day." in text and EXCLUDED not in text
    # The request for the newer sources is a new revision of the same presentation.
    monkeypatch.setattr(presentation_generation, "_dispatch_task", dispatch_now)
    renewed = ok(generate_v2(client, opportunity_id))
    assert renewed["presentation_id"] == presentation_id and renewed["presentation_version_id"] != str(v2_row["id"])
    assert stored_versions(presentation_id)[-1]["generation_source_manifest"]["approved_discovery_version_id"] == newer["id"]


def test_a_job_whose_frozen_sources_were_altered_fails_and_never_publishes(client: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    base = prepared(client)
    opportunity_id, presentation_id = base["opportunity"], base["presentation"]
    confirm(client, opportunity_id)
    store = get_memory_store()
    dispatched: list[tuple] = []
    monkeypatch.setattr(presentation_generation, "_dispatch_task", lambda task, *args: dispatched.append((task, args)))
    queued = ok(generate_v2(client, opportunity_id))
    enqueue = store.generation_jobs[UUID(queued["job_id"])]["result_json"]["_enqueue"]
    enqueue["master_v2_snapshot"]["findings"]["decisions"].append({"text": "The budget of 2 million is approved.", "source": "transcript"})

    task, args = dispatched[0]
    with pytest.raises(Exception, match="frozen source snapshot does not match"):
        task(*args)
    job = store.generation_jobs[UUID(queued["job_id"])]
    assert (job["status"], job["error_code"]) == ("FAILED", "MASTER_V2_SNAPSHOT_INVALID")
    assert len(stored_versions(presentation_id)) == 1, "no version is created from altered sources"
    status = v2_status(client, opportunity_id)
    assert (status["state"], status["job"]["error_code"], status["can_generate"]) == ("failed", "MASTER_V2_SNAPSHOT_INVALID", True)


def test_a_failed_generation_keeps_every_ready_version_and_can_be_repeated(client: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    base = prepared(client)
    opportunity_id, presentation_id = base["opportunity"], base["presentation"]
    confirm(client, opportunity_id)
    (v1_row,) = stored_versions(presentation_id)
    v1_pptx = download(client, presentation_id, str(v1_row["id"]))
    original = office.render_pptx

    def broken(*_args, **_kwargs):
        raise office.OfficeRenderError("LibreOffice produced no PDF for the deck")

    monkeypatch.setattr(office, "render_pptx", broken)
    # In-process execution reports the failure on the request; with a worker it is the job that fails.
    failed = generate_v2(client, opportunity_id)
    assert failed.status_code == 422 and "MASTER_PRESENTATION_RENDER_FAILED" in failed.text, failed.text
    status = v2_status(client, opportunity_id)
    assert (status["state"], status["latest_ready"], status["can_generate"]) == ("failed", None, True)
    assert status["job"]["error_code"] == "MASTER_PRESENTATION_RENDER_FAILED"
    # V1 is still there, ready and identical; the failed attempt published nothing.
    assert download(client, presentation_id, str(v1_row["id"])) == v1_pptx
    assert [(item["version_number"], item["source"]["product_version"]) for item in versions(client, presentation_id)] == [(1, "V1")]
    workflow = ok(get(client, f"/opportunities/{opportunity_id}/workflow-status"))
    assert workflow["documents"]["ppt1"]["latest_ready_version_id"] == str(v1_row["id"])
    assert workflow["documents"]["ppt2"]["latest_ready_version_id"] is None
    assert review(client, opportunity_id)["readiness"]["ready_for_v2"] is True

    monkeypatch.setattr(office, "render_pptx", original)
    repeated = ok(generate_v2(client, opportunity_id))
    assert repeated["status"] == "COMPLETED" and repeated["presentation_id"] == presentation_id
    ready = [row for row in stored_versions(presentation_id) if row["status"] == "ready"]
    assert [row["generation_source_manifest"]["product_version"] for row in ready] == ["V1", "V2"]
    assert v2_status(client, opportunity_id)["state"] == "ready"
    assert len(get_memory_store().presentations) == 1


def test_owner_review_and_finalization_pin_the_reviewed_v2_version(client: TestClient) -> None:
    base = prepared(client)
    opportunity_id, presentation_id = base["opportunity"], base["presentation"]
    step = lambda key: {s["key"]: s["state"] for s in ok(get(client, f"/opportunities/{opportunity_id}/workflow-status"))["steps"]}[key]  # noqa: E731

    confirm(client, opportunity_id, {"follow_ups": [EXCLUDED]})
    # Confirmed findings are not an owner review, and there is nothing to review yet.
    assert step("owner_review") == "pending"
    early = post(client, f"/opportunities/{opportunity_id}/workflow/owner-reviewed")
    assert early.status_code == 400 and "PPT2_NOT_GENERATED" in early.text
    assert post(client, f"/opportunities/{opportunity_id}/workflow/finalize").status_code == 400

    first = ok(generate_v2(client, opportunity_id))
    assert step("ppt2_generated") == "completed" and step("owner_review") != "completed"
    not_reviewed = post(client, f"/opportunities/{opportunity_id}/workflow/finalize")
    assert not_reviewed.status_code == 400 and "OWNER_REVIEW_REQUIRED" in not_reviewed.text
    ok(post(client, f"/opportunities/{opportunity_id}/workflow/owner-reviewed"))
    store = get_memory_store()
    opportunity = store.opportunities[UUID(opportunity_id)]
    assert str(opportunity["owner_reviewed_presentation_version_id"]) == first["presentation_version_id"]
    assert step("owner_review") == "completed"
    reviewed_at = opportunity["owner_reviewed_at"]
    ok(post(client, f"/opportunities/{opportunity_id}/workflow/owner-reviewed"))
    assert opportunity["owner_reviewed_at"] == reviewed_at, "repeating the review for the same version changes nothing"

    # A newer V2 revision has not been reviewed: the review no longer counts and finalizing is refused.
    confirm(client, opportunity_id)
    second = ok(generate_v2(client, opportunity_id))
    assert step("owner_review") != "completed"
    outdated = post(client, f"/opportunities/{opportunity_id}/workflow/finalize")
    assert outdated.status_code == 400 and "OWNER_REVIEW_OUTDATED" in outdated.text
    assert ok(get(client, f"/opportunities/{opportunity_id}/workflow-status"))["finalization"] is None
    ok(post(client, f"/opportunities/{opportunity_id}/workflow/owner-reviewed"))
    assert str(opportunity["owner_reviewed_presentation_version_id"]) == second["presentation_version_id"]

    finalized = ok(post(client, f"/opportunities/{opportunity_id}/workflow/finalize"))
    package = finalized["finalization"]
    v1_row = stored_versions(presentation_id)[0]
    assert package["ppt2_presentation_id"] == package["ppt1_presentation_id"] == presentation_id
    assert package["ppt2_version_id"] == second["presentation_version_id"]
    assert package["ppt1_version_id"] == str(v1_row["id"])
    assert package["approved_discovery_version_id"] == base["approved"]
    manifest = package["ppt2_generation_source_manifest"]
    assert (manifest["kind"], manifest["product_version"]) == ("master_presentation_v2", "V2")
    assert manifest == stored_versions(presentation_id)[-1]["generation_source_manifest"]
    assert {s["key"]: s["state"] for s in finalized["steps"]}["finalized"] == "completed"
    assert finalized["documents"]["ppt2"]["latest_ready_version_id"] == second["presentation_version_id"]

    # Finalized: no new V2, and later changes elsewhere do not replace the pinned package.
    locked = generate_v2(client, opportunity_id)
    assert locked.status_code == 409 and "WORKFLOW_FINALIZED" in locked.text
    assert v2_status(client, opportunity_id)["can_generate"] is False
    edited = client.patch(
        f"/opportunities/{opportunity_id}/discovery-paper",
        headers=headers(),
        json={"edits": [{"target": "thesis", "value": {"text": "Working hypothesis: approved after finalization"}}]},
    )
    assert edited.status_code == 200, edited.text
    ok(post(client, f"/opportunities/{opportunity_id}/discovery-paper/approve"))
    regenerated = ok(post(client, f"/opportunities/{opportunity_id}/stage1-outputs/presentation/regenerate"))["outputs"]["presentation"]
    assert regenerated["presentation_id"] == presentation_id
    later = ok(get(client, f"/opportunities/{opportunity_id}/workflow-status"))
    assert later["finalization"] == package
    assert later["documents"]["ppt2"]["latest_ready_version_id"] == second["presentation_version_id"]
    assert later["documents"]["ppt1"]["latest_ready_version_id"] == str(v1_row["id"])
    assert ok(post(client, f"/opportunities/{opportunity_id}/workflow/finalize"))["finalization"] == package
    assert [item["source"]["product_version"] for item in versions(client, presentation_id)] == ["V1", "V2", "V2", "V1"]


def test_a_new_pre_meeting_revision_after_v2_stays_v1_and_leaves_v2_in_place(client: TestClient) -> None:
    base = prepared(client)
    opportunity_id, presentation_id = base["opportunity"], base["presentation"]
    confirm(client, opportunity_id)
    v2 = ok(generate_v2(client, opportunity_id))
    # The same approved Discovery: nothing to regenerate, although the latest version is a V2.
    same = ok(post(client, f"/opportunities/{opportunity_id}/stage1-outputs/presentation/regenerate"))["outputs"]["presentation"]
    assert (same["presentation_id"], same["status"]) == (presentation_id, "ready")
    assert len(stored_versions(presentation_id)) == 2

    edited = client.patch(
        f"/opportunities/{opportunity_id}/discovery-paper",
        headers=headers(),
        json={"edits": [{"target": "thesis", "value": {"text": "Working hypothesis: a second approved thesis"}}]},
    )
    assert edited.status_code == 200, edited.text
    ok(post(client, f"/opportunities/{opportunity_id}/discovery-paper/approve"))
    ok(post(client, f"/opportunities/{opportunity_id}/stage1-outputs/presentation/regenerate"))
    rows = stored_versions(presentation_id)
    assert [(row["version_number"], row["generation_source_manifest"]["product_version"], row["journey_stage"]) for row in rows] == [
        (1, "V1", "first_contact"), (2, "V2", "post_meeting"), (3, "V1", "first_contact"),
    ]
    workflow = ok(get(client, f"/opportunities/{opportunity_id}/workflow-status"))
    assert workflow["documents"]["ppt1"]["latest_ready_version_id"] == str(rows[2]["id"])
    assert workflow["documents"]["ppt2"]["latest_ready_version_id"] == v2["presentation_version_id"]
    assert ok(get(client, f"/presentations/{presentation_id}/versions/{v2['presentation_version_id']}/deck"))["source"]["product_version"] == "V2"
    # The meeting information now has to be confirmed for the new approved Discovery before another V2.
    after = review(client, opportunity_id)
    assert after["master_presentation"]["version_id"] == str(rows[2]["id"])
    assert after["confirmation"]["stale_reasons"] == ["APPROVED_DISCOVERY_CHANGED"]
    assert generate_v2(client, opportunity_id).status_code == 409


def test_v2_never_contains_another_opportunitys_meeting(client: TestClient) -> None:
    first = prepared(client)
    other_meeting = "Cara: Requirement: Warehouse scans must post in real time.\nCara: Decision: Start in the Hamburg depot.\n"
    second = prepared(client, OTHER, body={**CLIENT, "client_name": "Südhafen Logistik"}, meeting=other_meeting, notes=None)
    view = ok(get(client, f"/opportunities/{second['opportunity']}/post-meeting-review", OTHER))
    ok(
        post(
            client,
            f"/opportunities/{second['opportunity']}/post-meeting-review/confirm",
            {
                "transcript_id": view["extraction"]["transcript_id"],
                "extraction_generated_at": view["extraction"]["generated_at"],
                "review_fingerprint": view["review_fingerprint"],
            },
            OTHER,
        )
    )
    confirm(client, first["opportunity"])
    mine = ok(generate_v2(client, first["opportunity"]))
    theirs = ok(generate_v2(client, second["opportunity"], OTHER))
    assert mine["presentation_id"] != theirs["presentation_id"]
    assert generate_v2(client, second["opportunity"]).status_code in (403, 404)
    my_text = " ".join(t for slide in slide_texts(download(client, first["presentation"], mine["presentation_version_id"])) for t in slide)
    their_pptx = client.get(
        f"/presentations/{second['presentation']}/versions/{theirs['presentation_version_id']}/download/pptx", headers=headers(OTHER)
    )
    their_text = " ".join(t for slide in slide_texts(their_pptx.content) for t in slide)
    assert "Warehouse scans" not in my_text and "Hamburg depot" not in my_text and "Südhafen" not in my_text
    assert "Quotes must go out" not in their_text and "Nordwind" not in their_text
    assert "Warehouse scans must post in real time." in their_text
    assert get(client, f"/presentations/{second['presentation']}/versions/{theirs['presentation_version_id']}/deck").status_code in (403, 404)


def test_competing_requests_create_one_job_and_one_version(client: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    """Simultaneous identical requests: the store's generation lock admits exactly one job."""
    import threading

    base = prepared(client)
    opportunity_id, presentation_id = base["opportunity"], base["presentation"]
    confirm(client, opportunity_id, {"follow_ups": [EXCLUDED]})
    store = get_memory_store()
    dispatched: list[tuple] = []
    guard = threading.Lock()

    def collect(task, *args) -> None:
        with guard:
            dispatched.append((task, args))

    monkeypatch.setattr(presentation_generation, "_dispatch_task", collect)
    token = headers()
    barrier = threading.Barrier(8)
    answers: list[tuple[int, dict]] = []

    def request() -> None:
        barrier.wait()
        response = client.post(f"/opportunities/{opportunity_id}/master-presentation/v2/generate", headers=token)
        with guard:
            answers.append((response.status_code, response.json()))

    threads = [threading.Thread(target=request) for _ in range(8)]
    for thread in threads:
        thread.start()
    for thread in threads:
        thread.join()

    assert [status for status, _body in answers] == [200] * 8, answers
    v2_jobs = [row for row in store.generation_jobs.values() if row.get("generation_lock_key") == "master_presentation_v2"]
    assert len(v2_jobs) == 1 and len(dispatched) == 1, "one job, dispatched once"
    assert {body["job_id"] for _status, body in answers} == {str(v2_jobs[0]["id"])}
    assert sorted(body["is_existing_job"] for _status, body in answers) == [False] + [True] * 7
    assert {body["presentation_id"] for _status, body in answers} == {presentation_id}
    assert v2_jobs[0]["generation_fingerprint"] == v2_jobs[0]["result_json"]["_enqueue"]["generation_source_manifest"]["generation_fingerprint"]
    # The presentation points at the plan of the job that won, so the worker accepts it.
    assert str(store.presentations[UUID(presentation_id)]["presentation_plan_id"]) == v2_jobs[0]["result_json"]["_enqueue"]["presentation_plan_id"]
    assert len(stored_versions(presentation_id)) == 1

    # Different confirmed information while that job is active is refused, not queued behind it.
    confirm(client, opportunity_id)
    other = generate_v2(client, opportunity_id)
    assert other.status_code == 409 and "PRESENTATION_GENERATION_IN_PROGRESS" in other.text
    assert len([row for row in store.generation_jobs.values() if row.get("generation_lock_key")]) == 1

    task, args = dispatched[0]
    task(*args)
    rows = stored_versions(presentation_id)
    assert [(row["version_number"], row["status"], row["generation_source_manifest"]["product_version"]) for row in rows] == [(1, "ready", "V1"), (2, "ready", "V2")]
    # The lock ended with the job: the newer confirmation can now be generated, as one more version.
    monkeypatch.setattr(presentation_generation, "_dispatch_task", lambda task, *args: task(*args))
    renewed = ok(generate_v2(client, opportunity_id))
    assert renewed["presentation_id"] == presentation_id and len(stored_versions(presentation_id)) == 3
    assert len(store.presentations) == 1


def test_the_generation_lock_holds_across_separate_store_instances(monkeypatch: pytest.MonkeyPatch) -> None:
    """Two API processes are two store instances over one database: the unique index decides."""
    import threading
    from pathlib import Path
    from uuid import uuid4

    import httpx

    from app.services import job_service
    from app.services.data.supabase_store import SupabaseDataStore

    table: list[dict] = []
    guard = threading.Lock()

    def request(self, method, name, *, json_body=None, params=None, headers=None):
        assert name == "generation_jobs"
        params = params or {}
        with guard:  # the database applies each statement atomically
            if method == "POST":
                row = {"id": str(uuid4()), "created_at": f"2026-10-08T12:00:{len(table):02d}+00:00", **json_body}
                # generation_jobs_one_active_locked_generation
                if row.get("generation_lock_key") and any(
                    other["opportunity_id"] == row["opportunity_id"]
                    and other.get("generation_lock_key") == row["generation_lock_key"]
                    and other["status"] in ("QUEUED", "RUNNING")
                    for other in table
                ):
                    return httpx.Response(409, json={"code": "23505", "message": "duplicate key value violates unique constraint"})
                table.append(row)
                return httpx.Response(201, json=[dict(row)])
            matched = [row for row in table if all(str(row.get(key)) == value[3:] for key, value in params.items() if value.startswith("eq."))]
            if method == "PATCH":
                for row in matched:
                    row.update(json_body)
            return httpx.Response(200, json=[dict(row) for row in matched])

    monkeypatch.setattr(SupabaseDataStore, "_request", request)
    stores = [SupabaseDataStore(f"token-of-process-{index}") for index in range(6)]
    assert len({id(store) for store in stores}) == 6
    opportunity_id, other_opportunity = uuid4(), uuid4()
    barrier = threading.Barrier(len(stores))
    outcomes: list[object] = []

    def compete(store: SupabaseDataStore) -> None:
        barrier.wait()
        try:
            outcome: object = job_service.create_job(
                opportunity_id=opportunity_id, job_type="presentation_generation", repository=store,
                enqueue={"generation_source_manifest": {"kind": "master_presentation_v2"}},
                generation_lock_key="master_presentation_v2", generation_fingerprint="a" * 64,
            )
        except job_service.GenerationLockConflict as exc:
            outcome = exc
        with guard:
            outcomes.append(outcome)

    threads = [threading.Thread(target=compete, args=(store,)) for store in stores]
    for thread in threads:
        thread.start()
    for thread in threads:
        thread.join()
    winners = [outcome for outcome in outcomes if isinstance(outcome, job_service.Job)]
    assert len(winners) == 1 and len(outcomes) - len(winners) == 5, "exactly one process gets the lock"
    assert len(table) == 1 and table[0]["generation_fingerprint"] == "a" * 64
    held = job_service.active_locked_job(stores[3], opportunity_id, "master_presentation_v2")
    assert held is not None and str(held["id"]) == str(winners[0].id), "every process sees the same holder"

    # Other opportunities and jobs without a lock key are not affected.
    job_service.create_job(opportunity_id=other_opportunity, job_type="presentation_generation", repository=stores[1],
                           generation_lock_key="master_presentation_v2", generation_fingerprint="b" * 64)
    job_service.create_job(opportunity_id=opportunity_id, job_type="presentation_generation", repository=stores[2])
    with pytest.raises(job_service.GenerationLockConflict):
        job_service.create_job(opportunity_id=opportunity_id, job_type="presentation_generation", repository=stores[4],
                               generation_lock_key="master_presentation_v2", generation_fingerprint="c" * 64)
    # The lock ends with the job, whether it completed or failed.
    for final in ("COMPLETED", "FAILED"):
        active = job_service.active_locked_job(stores[0], opportunity_id, "master_presentation_v2")
        stores[5].update_generation_job(UUID(str(active["id"])), {"status": final})
        assert job_service.active_locked_job(stores[0], opportunity_id, "master_presentation_v2") is None
        job_service.create_job(opportunity_id=opportunity_id, job_type="presentation_generation", repository=stores[0],
                               generation_lock_key="master_presentation_v2", generation_fingerprint="d" * 64)
    migration = (Path(__file__).resolve().parents[3] / "apps/services/api/supabase/migrations/042_master_presentation_v2.sql").read_text(encoding="utf-8")
    assert "CREATE UNIQUE INDEX IF NOT EXISTS generation_jobs_one_active_locked_generation" in migration
    assert "WHERE generation_lock_key IS NOT NULL AND status IN ('QUEUED', 'RUNNING')" in migration
    assert "ADD COLUMN IF NOT EXISTS owner_reviewed_presentation_version_id UUID" in migration
    assert "CREATE TABLE" not in migration and "CREATE POLICY" not in migration.upper()
